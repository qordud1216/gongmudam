"""기관의 표 형태 XLSX 서식: 명시적으로 지정한 시트·영역만 취합한다."""
import io
import zipfile
from datetime import date,datetime
from typing import Literal
from pydantic import BaseModel,ConfigDict,Field,model_validator
from openpyxl import load_workbook


def text(value):
    if value is None:return ''
    if isinstance(value,datetime):return value.date().isoformat()
    if isinstance(value,date):return value.isoformat()
    return str(value).strip()


class Registration(BaseModel):
    model_config=ConfigDict(extra='forbid',str_strip_whitespace=True)
    institution:str=Field(min_length=2,max_length=100)
    name:str=Field(min_length=2,max_length=100)
    version_label:str=Field(min_length=1,max_length=40)
    sheet_name:str=Field(min_length=1,max_length=31)
    header_row:int=Field(ge=1,le=50)
    header_end:int|None=Field(default=None,ge=1,le=50)
    data_start:int=Field(ge=2,le=60)
    data_end:int|None=Field(default=None,ge=2,le=2060)
    kind:Literal['generic','education']='generic'
    roles:dict[str,int]=Field(default_factory=dict,max_length=6)
    column_types:dict[int,Literal['text','number','date']]=Field(default_factory=dict,max_length=40)
    required:list[int]=Field(default_factory=list,max_length=40)
    duplicate_keys:list[int]=Field(default_factory=list,max_length=40)
    fill_merged_columns:list[int]=Field(default_factory=list,max_length=40)
    record_rows:int=Field(default=1,ge=1,le=20)
    group_columns:dict[int,str]=Field(default_factory=dict,max_length=40)
    positive:str=Field(default='희망',min_length=1,max_length=40)
    negative:str=Field(default='미희망',min_length=1,max_length=40)
    selected_area_confirmed:bool=False

    @model_validator(mode='after')
    def valid(self):
        end=self.header_end or self.header_row
        if not self.header_row<=end<self.header_row+10:raise ValueError('제목은 시작 행부터 최대 10개 행까지 지정하세요.')
        if self.data_start<=end:raise ValueError('데이터 시작 행은 제목 끝 행 다음이어야 합니다.')
        if self.data_end is not None and not self.data_start<=self.data_end<self.data_start+2000:raise ValueError('데이터 끝 행은 시작 행 이후 최대 2,000행 이내로 지정하세요.')
        if not self.selected_area_confirmed:raise ValueError('선택한 시트·영역만 취합함을 확인하세요.')
        if not set(self.roles).issubset({'department','name','rank','course','date','wish'}):raise ValueError('지원하지 않는 열 의미입니다.')
        if len(set(self.roles.values()))!=len(self.roles):raise ValueError('한 열에 두 가지 의미를 지정할 수 없습니다.')
        if self.kind=='education' and not {'department','name','course','date','wish'}.issubset(self.roles):raise ValueError('교육 명단은 부서·성명·교육명·교육일·참석 희망 열을 연결하세요.')
        if self.positive==self.negative:raise ValueError('희망·미희망 표시는 서로 달라야 합니다.')
        if self.record_rows>1 and (self.data_end is None or (self.data_end-self.data_start+1)%self.record_rows):
            raise ValueError('행 묶음은 데이터 끝 행을 지정하고 전체 행 수를 한 건의 행 수 배수로 맞추세요.')
        for mode in self.group_columns.values():
            if mode not in ('consistent','join') and mode not in [f'row:{i}' for i in range(self.record_rows)]:
                raise ValueError('행 묶음 읽기 방식 또는 묶음 안 행 번호가 올바르지 않습니다.')
        if self.record_rows==1 and self.group_columns:raise ValueError('여러 행 묶음에서만 열별 묶음 기준을 지정하세요.')
        return self


def open_book(raw):
    try:
        if len(raw)>2*1024*1024:raise ValueError('2MB 이하 XLSX만 지원합니다.')
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            if len(z.infolist())>256 or sum(f.file_size for f in z.infolist())>20*1024*1024:raise ValueError('압축 해제 크기 또는 항목 수 제한 초과')
            if any('vbaProject' in f.filename or 'externalLinks/' in f.filename for f in z.infolist()):raise ValueError('매크로·외부 연결 파일은 지원하지 않습니다.')
        wb=load_workbook(io.BytesIO(raw),data_only=False,keep_links=False)
        if len(wb.worksheets)>20:wb.close();raise ValueError('시트는 최대 20개까지 지원합니다.')
        return wb
    except ValueError:raise
    except Exception as exc:raise ValueError('읽을 수 없는 XLSX입니다. 파일 형식·암호·손상 여부를 확인하세요.') from exc


def inspect(raw):
    wb=open_book(raw)
    try:return {'sheets':[{'name':ws.title,'rows':ws.max_row,'columns':ws.max_column,
          'merged_ranges':[str(r) for r in ws.merged_cells.ranges],
          'preview':[[text(v)[:160] for v in row] for row in ws.iter_rows(min_row=1,max_row=min(50,ws.max_row),max_col=min(50,ws.max_column),values_only=True)]} for ws in wb]}
    finally:wb.close()


def header_layout(ws,first,last):
    if not 1<=first<=last<=50 or last-first>=10:raise ValueError('제목 범위는 1~50행 중 최대 10개 행으로 지정하세요.')
    if ws.max_column>50:raise ValueError('시작 여백을 포함해 50열 이내의 표를 지원합니다.')
    merges=[m for m in ws.merged_cells.ranges if m.min_row<=last and m.max_row>=first]
    merged_owner={}
    for m in merges:
        if m.min_row<first or m.max_row>last:raise ValueError(f'제목 병합 {m} 전체를 포함하도록 제목 시작·끝 행을 지정하세요.')
        for r in range(m.min_row,m.max_row+1):
            for col in range(m.min_col,m.max_col+1):
                if (r,col) in merged_owner:raise ValueError('제목 병합 범위가 서로 겹칩니다. 원본 양식을 확인하세요.')
                merged_owner[r,col]=m
    effective={};occupied=[]
    for r in range(first,last+1):
        for col in range(1,ws.max_column+1):
            merge=merged_owner.get((r,col))
            cell=ws.cell(merge.min_row,merge.min_col) if merge else ws.cell(r,col)
            if cell.data_type=='f':raise ValueError('제목 영역에는 수식을 사용할 수 없습니다.')
            value=' '.join(text(cell.value).split())
            if value:occupied.append(col)
            effective[r,col]=(cell.coordinate,value)
    if not occupied:raise ValueError('선택한 범위에 열 제목이 없습니다.')
    start,end=min(occupied),max(occupied)
    paths=[];spans=[];previous_tokens=None
    for col in range(start,end+1):
        path=[];seen=set();tokens=[]
        for r in range(first,last+1):
            anchor,value=effective[r,col]
            if value and anchor not in seen:
                tokens.append((anchor,value))
                if not path or path[-1]!=value:path.append(value)
                seen.add(anchor)
        if tokens and tokens==previous_tokens:
            spans[-1]['end']=col
        else:
            paths.append(path);spans.append({'start':col,'end':col})
        previous_tokens=tokens
    headers=[' / '.join(path) for path in paths]
    if len(headers)>40 or any(not h or len(h)>240 for h in headers) or len(set(headers))!=len(headers):
        raise ValueError('펼친 열 제목에 빈칸·중복이 있습니다. 하위 제목까지 포함해 범위를 지정하세요(열당 240자 이내).')
    matrix=[[text(ws.cell(r,col).value) for col in range(start,end+1)] for r in range(first,last+1)]
    return {'header_schema_version':3,'header_row':first,'header_end':last,'headers':headers,'header_paths':paths,'column_spans':spans,
            'header_cells':matrix,'header_merges':sorted(str(m) for m in merges), 'start_column':start}


def preview(raw,sheet_name,first,last):
    wb=open_book(raw)
    try:
        if sheet_name not in wb.sheetnames:raise ValueError('선택한 시트가 없습니다.')
        return header_layout(wb[sheet_name],first,last)
    finally:wb.close()


def body_merges(ws,snapshot,last):
    spans=logical_spans(snapshot);start=spans[0]['start'];end=spans[-1]['end']
    allowed=set(snapshot.get('fill_merged_columns',[]));result=[]
    for m in ws.merged_cells.ranges:
        if m.max_row<snapshot['data_start'] or m.min_row>last or m.max_col<start or m.min_col>end:continue
        if m.min_row<snapshot['data_start'] or m.max_row>last or m.min_col<start or m.max_col>end:
            raise ValueError(f'병합 {m}이 데이터 영역 경계를 넘습니다. 범위를 다시 지정하세요.')
        index=next((i for i,span in enumerate(spans) if span['start']==m.min_col and span['end']==m.max_col),None)
        if index is None:raise ValueError(f'데이터의 병합 {m}이 제목의 서로 다른 항목을 합치거나 한 항목의 일부만 병합합니다. 제목과 같은 항목 범위를 사용하세요.')
        if m.max_row>m.min_row and index not in allowed:raise ValueError(f'{m}의 열은 세로 병합값 펼치기를 허용하지 않았습니다. 등록 설정을 확인하세요.')
        size=snapshot.get('record_rows',1)
        if size>1 and (m.min_row-snapshot['data_start'])//size!=(m.max_row-snapshot['data_start'])//size:
            raise ValueError(f'병합 {m}이 서로 다른 행 묶음을 넘습니다. 한 건 안에서만 병합하세요.')
        if ws.cell(m.min_row,m.min_col).data_type=='f':raise ValueError(f'병합 원본 {m}에 수식이 있습니다. 값으로 회신하세요.')
        result.append(m)
    return result


def logical_spans(snapshot):
    return snapshot.get('column_spans') or [{'start':snapshot['start_column']+i,'end':snapshot['start_column']+i} for i in range(len(snapshot['headers']))]


def describe(raw,config):
    wb=open_book(raw)
    try:
        if config.sheet_name not in wb.sheetnames:raise ValueError('선택한 시트가 없습니다.')
        ws=wb[config.sheet_name]
        structure=header_layout(ws,config.header_row,config.header_end or config.header_row)
        headers=structure['headers']
        indexes=list(config.roles.values())+config.required+config.duplicate_keys+list(config.column_types)+config.fill_merged_columns+list(config.group_columns)
        if any(i<0 or i>=len(headers) for i in indexes):raise ValueError('열 연결 또는 확인 기준이 실제 열 범위를 벗어났습니다.')
        required=sorted(set(config.required)|set(config.roles.values() if config.kind=='education' else [] )-({config.roles['rank']} if 'rank' in config.roles else set()))
        types={str(i):kind for i,kind in config.column_types.items()}
        if 'date' in config.roles:types.setdefault(str(config.roles['date']),'date')
        for i,mode in config.group_columns.items():
            if mode=='join' and (types.get(str(i),'text')!='text' or i in config.roles.values() or i in config.duplicate_keys):
                raise ValueError('여러 값 합치기는 일반 문자 열에만 지정하세요. 역할·중복 기준·숫자·날짜 열에는 공통값 확인이나 특정 행을 사용하세요.')
        spec=config.model_dump()|structure|{'required':required,'column_types':types,
              'fill_merged_columns':sorted(set(config.fill_merged_columns)),
              'duplicate_keys':list(dict.fromkeys(config.duplicate_keys)),
              'ignored_sheets':[s.title for s in wb if s.title!=ws.title]}
        last=config.data_end if config.record_rows>1 else min(ws.max_row,config.data_end) if config.data_end else ws.max_row
        body_merges(ws,spec,last)
        return spec
    finally:wb.close()


def parse_legacy(raw,snapshot):
    wb=open_book(raw)
    try:
        if snapshot['sheet_name'] not in wb.sheetnames:raise ValueError('등록 서식과 시트명이 다릅니다.')
        ws=wb[snapshot['sheet_name']];start=snapshot['start_column'];width=len(snapshot['headers']);end=start+width-1
        if ws.max_column>50:raise ValueError('등록 서식 범위를 벗어나는 열이 있습니다.')
        actual=[text(ws.cell(snapshot['header_row'],i).value) for i in range(start,end+1)]
        if actual!=snapshot['headers']:raise ValueError('등록 서식과 제목 행의 위치·이름·열 순서가 다릅니다. 다른 양식은 별도 서식으로 등록하세요.')
        if any(text(ws.cell(snapshot['header_row'],i).value) for i in range(1,ws.max_column+1) if i<start or i>end):raise ValueError('등록 서식에 없는 열 제목이 추가되었습니다.')
        last=min(ws.max_row,snapshot['data_end']) if snapshot['data_end'] else ws.max_row
        if last>=snapshot['data_start']+2000:raise ValueError('파일당 데이터 2,000행까지 지원합니다. 합계·꼬리말은 데이터 끝 행으로 제외하세요.')
        for merged in ws.merged_cells.ranges:
            touches_header=merged.min_row<=snapshot['header_row']<=merged.max_row
            touches_data=merged.max_row>=snapshot['data_start'] and merged.min_row<=last
            if (touches_header or touches_data) and merged.min_col<=end and merged.max_col>=start:raise ValueError('제목 행과 데이터 영역의 셀 병합은 지원하지 않습니다.')
        result=[]
        for row in ws.iter_rows(min_row=snapshot['data_start'],max_row=last,max_col=max(end,ws.max_column)):
            if any(text(c.value) for c in row if c.column<start or c.column>end):raise ValueError(f'{row[0].row}행에 등록 영역 밖 데이터가 있습니다.')
            cells=row[start-1:end]
            if any(c.data_type=='f' for c in cells):raise ValueError(f'{row[0].row}행에 수식이 있습니다. 값으로 붙여넣어 회신하세요.')
            values=[text(c.value) for c in cells]
            if values==snapshot['headers']:raise ValueError(f'{row[0].row}행에 제목 행이 반복되었습니다. 중간 제목을 제거하세요.')
            if any(values):result.append((row[0].row,values))
        if not result:raise ValueError('등록 영역에 데이터가 없습니다. 자료 없음 회신을 별도로 기록하세요.')
        return result
    finally:wb.close()


def parse(raw,snapshot):
    if snapshot.get('header_schema_version',1)<2:return parse_legacy(raw,snapshot)
    wb=open_book(raw)
    try:
        if snapshot['sheet_name'] not in wb.sheetnames:raise ValueError('등록 서식과 시트명이 다릅니다.')
        ws=wb[snapshot['sheet_name']]
        actual=header_layout(ws,snapshot['header_row'],snapshot['header_end'])
        for key in ('headers','header_cells','header_merges','start_column')+ (('column_spans',) if 'column_spans' in snapshot else ()):
            if actual[key]!=snapshot[key]:raise ValueError('등록 서식과 제목 값·열 위치·병합 구조가 다릅니다. 같은 양식인지 확인하세요.')
        spans=logical_spans(snapshot);start=spans[0]['start'];end=spans[-1]['end']
        grouped=snapshot.get('record_rows',1)>1
        last=snapshot['data_end'] if grouped else min(ws.max_row,snapshot['data_end']) if snapshot['data_end'] else ws.max_row
        if last>=snapshot['data_start']+2000:raise ValueError('파일당 데이터 2,000행까지 지원합니다.')
        merges=body_merges(ws,snapshot,last);result=[];merge_lookup={}
        span_indexes={(s['start'],s['end']):i for i,s in enumerate(spans)}
        for m in merges:
            i=span_indexes[m.min_col,m.max_col]
            for row_number in range(m.min_row,m.max_row+1):
                if (row_number,i) in merge_lookup:raise ValueError('데이터 병합 범위가 서로 겹칩니다. 원본 양식을 확인하세요.')
                merge_lookup[row_number,i]=m
        for row in ws.iter_rows(min_row=snapshot['data_start'],max_row=last,max_col=max(end,ws.max_column)):
            if any(text(c.value) for c in row if c.column<start or c.column>end):raise ValueError(f'{row[0].row}행에 등록 영역 밖 데이터가 있습니다.')
            cells=row[start-1:end]
            if any(c.data_type=='f' for c in cells):raise ValueError(f'{row[0].row}행에 수식이 있습니다. 값으로 붙여넣어 회신하세요.')
            # 병합값만 내려오는 빈 행은 별도 레코드로 만들지 않는다.
            if not any(text(c.value) for c in cells) and (not grouped or not any((row[0].row,i) in merge_lookup for i in range(len(spans)))):continue
            values=[];provenance=[]
            for index,span in enumerate(spans):
                merge=merge_lookup.get((row[0].row,index))
                if span['start']!=span['end'] and not merge:raise ValueError(f'{row[0].row}행의 {snapshot["headers"][index]} 병합 범위가 등록한 한 항목의 열 범위와 다릅니다.')
                values.append(text(ws.cell(row[0].row,span['start']).value))
            for index in range(len(spans)):
                m=merge_lookup.get((row[0].row,index))
                if not m:continue
                anchor=ws.cell(m.min_row,m.min_col)
                values[index]=text(anchor.value)
                provenance.append({'column':index,'header':snapshot['headers'][index],'range':str(m),'anchor':anchor.coordinate,
                                   'value':values[index],'inherited':row[0].row!=m.min_row,
                                   'direction':'horizontal' if m.min_row==m.max_row else 'vertical' if m.min_col==m.max_col else 'block'})
            leaf_headers=[path[-1] for path in snapshot['header_paths']]
            if [' '.join(v.split()) for v in values] in (snapshot['headers'],leaf_headers):raise ValueError(f'{row[0].row}행에 제목이 반복되었습니다.')
            result.append((row[0].row,values,provenance))
        if grouped:result=group_records(ws,snapshot,result)
        if not result:raise ValueError('등록 영역에 데이터가 없습니다. 자료 없음 회신을 별도로 기록하세요.')
        return result
    finally:wb.close()


def group_records(ws,snapshot,physical):
    """고정 행 묶음: 원본의 모든 비어 있지 않은 셀과 적용 기준을 별도로 보존한다."""
    size=snapshot['record_rows'];by_row={r[0]:r for r in physical};result=[];spans=logical_spans(snapshot)
    for first in range(snapshot['data_start'],snapshot['data_end']+1,size):
        last=first+size-1;members=[by_row[r] for r in range(first,last+1) if r in by_row]
        if not any(text(ws.cell(r,c).value) for r in range(first,last+1) for c in range(spans[0]['start'],spans[-1]['end']+1)):continue
        values=[];merges=[];fields=[]
        for index,header in enumerate(snapshot['headers']):
            mode=snapshot.get('group_columns',{}).get(str(index),'consistent')
            entries=[(r,v[index]) for r,v,_ in members if v[index]]
            unique=list(dict.fromkeys(v for _,v in entries))
            if mode=='consistent':
                if len(unique)>1:raise ValueError(f'{first}~{last}행 묶음의 {header} 공통값이 다릅니다. 사람·건이 섞였는지 또는 열별 읽기 기준을 확인하세요.')
                value=unique[0] if unique else ''
            elif mode=='join':value='\n'.join(v for _,v in entries)
            else:
                selected=first+int(mode.split(':')[1]);value=by_row.get(selected,(None,['']*len(spans)))[1][index]
            values.append(value)
            cells=[{'cell':ws.cell(r,span_col).coordinate,'value':text(ws.cell(r,span_col).value)}
                   for r in range(first,last+1) for span_col in range(spans[index]['start'],spans[index]['end']+1)
                   if text(ws.cell(r,span_col).value)]
            fields.append({'column':index,'header':header,'mode':mode,'cells':cells,'value':value})
            for _,_,provenance in members:
                for p in provenance:
                    if p['column']==index and p not in merges:merges.append(p)
        result.append((first,values,merges,{'start':first,'end':last,'record_rows':size,'fields':fields}))
    return result
