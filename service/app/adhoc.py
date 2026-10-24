"""로컬 단일 사용자용 수시업무·동일 서식 교육명단 취합. 모델 호출 없음."""
from contextlib import contextmanager
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Literal
import hashlib
import io
import json
import math
from decimal import Decimal,InvalidOperation
import sqlite3
import uuid
import zipfile

from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel, ConfigDict, Field, model_validator
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from . import adhoc_templates as templates
from . import adhoc_workflow as workflow
from . import adhoc_workspace

DB_PATH=Path(__file__).resolve().parent.parent/'data'/'adhoc.db'
router=APIRouter(prefix='/api/adhoc',tags=['수시업무·자료 취합'])
HEADERS=['부서','성명','직급','교육명','교육일','참석 희망']
MAX_BYTES=2*1024*1024


class Strict(BaseModel):
    model_config=ConfigDict(extra='forbid',str_strip_whitespace=True)


class Create(Strict):
    template_id: str|None=Field(default=None,max_length=40)
    situation: str=Field(min_length=5,max_length=3000)
    title: str=Field(min_length=2,max_length=150)
    course: str=Field(min_length=2,max_length=100)
    event_date: date
    deadline: date
    departments: list[str]=Field(min_length=1,max_length=30)

    @model_validator(mode='after')
    def valid(self):
        self.departments=[s.strip() for s in self.departments]
        if any(not s or len(s)>60 for s in self.departments) or len(set(self.departments))!=len(self.departments):
            raise ValueError('취합 대상 부서는 중복 없이 1~60자로 입력하세요.')
        if not self.template_id and self.deadline>self.event_date: raise ValueError('회신 마감일은 교육일 이후일 수 없습니다.')
        return self


class Version(Strict):
    version: int=Field(ge=1)


class Review(Version):
    values: list[str]=Field(min_length=1,max_length=40)
    decision: Literal['include','exclude']
    reason: str=Field(min_length=5,max_length=500)


class BatchReview(Version):
    reason: str=Field(min_length=5,max_length=500)


class EmptyReply(Version):
    department: str=Field(min_length=1,max_length=60)
    reason: str=Field(min_length=5,max_length=500)


def stamp():return datetime.now(timezone.utc).isoformat()
def encode(obj):return json.dumps(obj,ensure_ascii=False)


@contextmanager
def db():
    DB_PATH.parent.mkdir(parents=True,exist_ok=True)
    c=sqlite3.connect(DB_PATH,timeout=15);c.row_factory=sqlite3.Row
    c.execute('PRAGMA foreign_keys=ON')
    try:
        with c:yield c
    finally:c.close()


def initialize():
    with db() as c:c.executescript('''
    CREATE TABLE IF NOT EXISTS jobs(id TEXT PRIMARY KEY,version INTEGER NOT NULL,settings TEXT NOT NULL,created_at TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS replies(id TEXT PRIMARY KEY,job_id TEXT NOT NULL REFERENCES jobs(id),department TEXT NOT NULL,
      filename TEXT NOT NULL,sha256 TEXT NOT NULL,original BLOB NOT NULL,created_at TEXT NOT NULL,UNIQUE(job_id,sha256));
    CREATE TABLE IF NOT EXISTS rows(id TEXT PRIMARY KEY,job_id TEXT NOT NULL REFERENCES jobs(id),reply_id TEXT NOT NULL REFERENCES replies(id),
      source_row INTEGER NOT NULL,original_values TEXT NOT NULL,values_json TEXT NOT NULL,decision TEXT NOT NULL,reason TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS row_merge_provenance(row_id TEXT PRIMARY KEY REFERENCES rows(id),provenance TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS row_group_provenance(row_id TEXT PRIMARY KEY REFERENCES rows(id),provenance TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS empty_replies(job_id TEXT NOT NULL REFERENCES jobs(id),department TEXT NOT NULL,reason TEXT NOT NULL,
      created_at TEXT NOT NULL,PRIMARY KEY(job_id,department));
    CREATE TABLE IF NOT EXISTS history(id INTEGER PRIMARY KEY,job_id TEXT NOT NULL REFERENCES jobs(id),kind TEXT NOT NULL,detail TEXT NOT NULL,created_at TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS exports(id TEXT PRIMARY KEY,job_id TEXT NOT NULL REFERENCES jobs(id),version INTEGER NOT NULL,
      original BLOB NOT NULL,sha256 TEXT NOT NULL,created_at TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS collection_templates(id TEXT PRIMARY KEY,institution TEXT NOT NULL,name TEXT NOT NULL,
      version_label TEXT NOT NULL,config TEXT NOT NULL,filename TEXT NOT NULL,sha256 TEXT NOT NULL,original BLOB NOT NULL,
      created_at TEXT NOT NULL,UNIQUE(institution,name,version_label));
    ''')
    with db() as c:c.executescript(workflow.SCHEMA)
    with db() as c:c.executescript('''CREATE TABLE IF NOT EXISTS workflow_chat_requests (
      request_id TEXT PRIMARY KEY,job_id TEXT NOT NULL REFERENCES jobs(id),query_hash TEXT NOT NULL,
      response TEXT NOT NULL,created_at TEXT NOT NULL);''')


def job(c,key):
    r=c.execute('SELECT * FROM jobs WHERE id=?',(key,)).fetchone()
    if not r:raise HTTPException(404,'수시업무를 찾을 수 없습니다.')
    return dict(r)|{'settings':json.loads(r['settings'])}


def bump(c,key,version):
    if c.execute('UPDATE jobs SET version=version+1 WHERE id=? AND version=?',(key,version)).rowcount!=1:
        raise HTTPException(409,'다른 변경이 있습니다. 업무를 다시 열어 주세요.')


def log(c,key,kind,detail):
    c.execute('INSERT INTO history(job_id,kind,detail,created_at) VALUES (?,?,?,?)',(key,kind,encode(detail),stamp()))


def text(value):
    if value is None:return ''
    if isinstance(value,datetime):return value.date().isoformat()
    if isinstance(value,date):return value.isoformat()
    return str(value).strip()


def layout(settings):
    return settings.get('template_snapshot') or {'kind':'education','headers':HEADERS,
      'roles':dict(department=0,name=1,rank=2,course=3,date=4,wish=5),
      'required':[0,1,3,4,5],'duplicate_keys':[0,1,3,4],'positive':'희망','negative':'미희망'}


def role_value(values,settings,role):
    index=layout(settings)['roles'].get(role)
    return values[index] if index is not None else ''


def attendee(values,settings):
    config=layout(settings)
    return config['kind']!='education' or role_value(values,settings,'wish')==config['positive']


def issues(values,settings,department):
    config=layout(settings)
    errors=[]
    if any(not values[i] for i in config['required']):errors.append('필수값 누락')
    if 'department' in config['roles'] and role_value(values,settings,'department')!=department:errors.append('회신 부서와 행의 부서 불일치')
    if config['kind']=='education':
        if role_value(values,settings,'course')!=settings['course']:errors.append('교육명 불일치')
        if role_value(values,settings,'date')!=settings['event_date']:errors.append('교육일 불일치 · YYYY-MM-DD 확인')
        if role_value(values,settings,'wish') not in (config['positive'],config['negative']):errors.append('참석 희망 표시가 등록한 값과 다릅니다.')
    if any(len(v)>200 for v in values):errors.append('셀 내용 200자 초과')
    if any(v.startswith(('=','+','-','@')) for i,v in enumerate(values) if config.get('column_types',{}).get(str(i))!='number'):errors.append('수식처럼 시작하는 값 확인 필요')
    for index,kind in config.get('column_types',{}).items():
        value=values[int(index)]
        if not value:continue
        try:
            if kind=='date':date.fromisoformat(value)
            if kind=='number':
                n=Decimal(value)
                if not n.is_finite() or len(n.as_tuple().digits)>15 or abs(n.adjusted())>100:raise ValueError()
        except (ValueError,InvalidOperation):errors.append(f"{config['headers'][int(index)]}: {'YYYY-MM-DD 날짜' if kind=='date' else '15자리 이내 숫자'} 형식 확인")
    return errors


def parse(raw):
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            if len(z.infolist())>256 or sum(f.file_size for f in z.infolist())>20*1024*1024:
                raise ValueError('압축 해제 크기 또는 항목 수 제한 초과')
            if any('vbaProject' in f.filename or 'externalLinks/' in f.filename for f in z.infolist()):
                raise ValueError('매크로·외부 연결 파일은 지원하지 않습니다.')
        wb=load_workbook(io.BytesIO(raw),data_only=False,keep_links=False)
        try:
            if '참석희망자' not in wb.sheetnames:raise ValueError('참석희망자 시트가 없습니다. 제공 서식을 사용하세요.')
            ws=wb['참석희망자']
            if ws.max_column!=6 or ws.max_row>2001 or ws.merged_cells:
                raise ValueError('6개 열·최대 2,000행의 병합 없는 서식만 지원합니다.')
            if [text(ws.cell(1,i).value) for i in range(1,7)]!=HEADERS:
                raise ValueError('제목 행이 제공 서식과 다릅니다. 열 이름·순서를 확인하세요.')
            if any(s.title not in ('참석희망자','작성안내') and s.max_row>0 and any(v is not None for row in s.values for v in row) for s in wb):
                raise ValueError('추가 데이터 시트가 있습니다. 취합할 자료는 참석희망자 시트에 모아 주세요.')
            rows=[]
            for row in ws.iter_rows(min_row=2):
                if any(cell.data_type=='f' for cell in row):raise ValueError(f'{row[0].row}행에 수식이 있습니다. 값으로 붙여넣은 파일을 사용하세요.')
                values=[text(cell.value) for cell in row]
                if any(values):rows.append((row[0].row,values))
            if not rows:raise ValueError('빈 회신 파일입니다. 희망자가 없다면 무응답으로 두지 말고 희망자 없음 회신을 기록하세요.')
            return rows
        finally:wb.close()
    except (ValueError,HTTPException):raise
    except Exception as exc:raise ValueError('읽을 수 없는 XLSX 파일입니다. 암호·손상·파일 형식을 확인하세요.') from exc


def _summarize_rows(rows,byfile,settings,empty):
    """Project reply/department totals without rescanning rows for every group."""
    file_counts={fid:{'rows':0,'pending':0} for fid in byfile}
    department_counts={}
    included=pending=excluded=clean_pending=attendees=0
    for row in rows:
        decision=row['decision']
        is_pending=decision=='pending'
        is_included=decision=='include'
        counts=file_counts[row['reply_id']]
        counts['rows']+=1;counts['pending']+=is_pending
        department=byfile[row['reply_id']]['department']
        counts=department_counts.setdefault(department,{'rows':0,'included':0,'pending':0,'active':0})
        counts['rows']+=1;counts['included']+=is_included
        counts['pending']+=is_pending;counts['active']+=decision!='exclude'
        included+=is_included;pending+=is_pending;excluded+=decision=='exclude'
        clean_pending+=is_pending and not row['issues']
        attendees+=is_included and attendee(row['values'],settings)
    statuses=[]
    for department in settings['departments']:
        counts=department_counts.get(department,{'rows':0,'included':0,'pending':0,'active':0})
        conflict=department in empty and counts['active']>0
        statuses.append({'department':department,
          'status':'회신 충돌' if conflict else '검토 필요' if counts['pending'] else '회신 수신' if counts['rows'] else '희망자 없음' if department in empty else '미회신',
          'rows':counts['rows'],'included':counts['included'],'empty_reply':empty.get(department),'conflict':conflict})
    for fid,file in byfile.items():file.update(file_counts[fid])
    return dict(departments=statuses,
      summary={'total':len(rows),'included':included,'pending':pending,'excluded':excluded,
               'missing':sum(s['status']=='미회신' for s in statuses),'clean_pending':clean_pending,'attendees':attendees},
      can_finalize=bool(byfile or empty) and pending==0 and not any(s['status'] in ('미회신','회신 충돌') for s in statuses))


def view(c,key):
    result=job(c,key);settings=result['settings']
    files=[dict(r) for r in c.execute('SELECT id,department,filename,sha256,created_at FROM replies WHERE job_id=? ORDER BY created_at,id',(key,))]
    byfile={f['id']:f for f in files}
    rows=[];groups={}
    merge_info={r['row_id']:json.loads(r['provenance']) for r in c.execute('SELECT p.* FROM row_merge_provenance p JOIN rows r ON r.id=p.row_id WHERE r.job_id=?',(key,))}
    group_info={r['row_id']:json.loads(r['provenance']) for r in c.execute('SELECT p.* FROM row_group_provenance p JOIN rows r ON r.id=p.row_id WHERE r.job_id=?',(key,))}
    duplicate_keys=layout(settings)['duplicate_keys']
    for r in c.execute('SELECT * FROM rows WHERE job_id=? ORDER BY rowid',(key,)):
        item=dict(r);item['values']=json.loads(item.pop('values_json'));item['original_values']=json.loads(item['original_values'])
        item['merge_provenance']=merge_info.get(item['id'],[])
        item['group_provenance']=group_info.get(item['id'])
        item['filename']=byfile[item['reply_id']]['filename'];item['issues']=issues(item['values'],settings,byfile[item['reply_id']]['department'])
        if item['decision']!='exclude' and duplicate_keys and all(item['values'][i] for i in duplicate_keys):
            groups.setdefault(tuple(item['values'][i] for i in duplicate_keys),[]).append(item)
        rows.append(item)
    for group in groups.values():
        if len(group)>1:
            for r in group:r['issues'].append('지정한 기준 열 중복 의심 · 원본 대조 필요')
    empty={r['department']:dict(r) for r in c.execute('SELECT department,reason,created_at FROM empty_replies WHERE job_id=?',(key,))}
    result.update(files=files,rows=rows,**_summarize_rows(rows,byfile,settings,empty),
      history=[dict(r) for r in c.execute('SELECT kind,detail,created_at FROM history WHERE job_id=? ORDER BY id DESC',(key,))],
      exports=[dict(r)|{'is_current':r['version']==result['version']} for r in c.execute('SELECT id,version,sha256,created_at FROM exports WHERE job_id=? ORDER BY created_at DESC',(key,))])
    result['request_draft']=f"[{settings['title']}] 참석 희망자 제출 요청\n교육명: {settings['course']}\n교육일: {settings['event_date']}\n회신 마감일: {settings['deadline']} (마감 시각 별도 확인)\n대상: {', '.join(settings['departments'])}\n첨부한 동일 서식의 제목 행과 시트명을 유지하여 회신해 주세요. 참석 희망자가 없는 경우에도 없음을 회신해 주세요.\n필수 항목: 부서, 성명, 교육명, 교육일, 참석 희망. 직급은 선택 항목입니다."
    result['layout']=layout(settings)
    if settings.get('template_snapshot'):
        config=settings['template_snapshot']
        result['request_draft']=f"[{settings['title']}] 자료 제출 요청\n자료명: {settings['course']}\n기준일: {settings['event_date']}\n회신 마감일: {settings['deadline']} (시각 별도 확인)\n대상: {', '.join(settings['departments'])}\n서식: {config['institution']} / {config['name']} / {config['version_label']}\n시트: {config['sheet_name']}, 제목 {config['header_row']}~{config.get('header_end') or config['header_row']}행, 데이터 {config['data_start']}행부터\n한 건의 행 수: {config.get('record_rows',1)}행. 행을 추가·삭제하지 말고 묶음 위치를 유지해 주세요.\n등록한 제목·병합 구조·열 순서를 유지해 회신해 주세요. 해당 자료가 없는 경우에도 없음을 회신해 주세요."
    result['workflow']=workflow.project(c,result)
    return result


def sheet(wb,name,headers,rows,column_types=None):
    ws=wb.create_sheet(name);ws.append(headers);ws.freeze_panes='A2';ws.sheet_view.showGridLines=False
    for row in rows:
        row=list(row)
        for index,kind in (column_types or {}).items():
            i=int(index)
            if not row[i]:continue
            try:
                if kind=='date':row[i]=date.fromisoformat(str(row[i]))
                if kind=='number':
                    n=Decimal(str(row[i]))
                    if n.is_finite() and len(n.as_tuple().digits)<=15 and abs(n.adjusted())<=100:row[i]=n
            except (ValueError,InvalidOperation):pass
        ws.append(row)
        # 사용자 입력을 수식으로 실행하지 않고 그대로 보존한다.
        for cell in ws[ws.max_row]:
            if isinstance(cell.value,str):cell.data_type='s';cell.number_format='@'
    for cell in ws[1]:cell.font=Font(name='맑은 고딕',bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='244C46');cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    from openpyxl.utils import get_column_letter
    for i,h in enumerate(headers,1):ws.column_dimensions[get_column_letter(i)].width=64 if h in ('병합 원본','원본 셀·값') else 50 if h in ('확인사항','처리 사유','읽기 기준') else 36 if h=='원본 파일' else 23
    ws.row_dimensions[1].height=max(26,min(400,max(math.ceil(sum(2 if ord(c)>255 else 1 for c in h)/(ws.column_dimensions[get_column_letter(i)].width-2))*17+8 for i,h in enumerate(headers,1))))
    for row in ws.iter_rows(min_row=2):
        height=24
        for cell in row:
            cell.font=Font(name='맑은 고딕',size=11);cell.alignment=Alignment(vertical='center',wrap_text=True)
            if isinstance(cell.value,date):cell.number_format='yyyy-mm-dd'
            if isinstance(cell.value,Decimal):cell.number_format='#,##0.###############'
            if headers[cell.column-1]=='교육일' and str(cell.column-1) not in (column_types or {}):
                try:cell.value=date.fromisoformat(str(cell.value));cell.number_format='yyyy-mm-dd'
                except ValueError:pass
            width=ws.column_dimensions[cell.column_letter].width
            lines=sum(max(1,math.ceil(sum(2 if ord(ch)>255 else 1 for ch in line)/max(1,width-2))) for line in str(cell.value or '').split('\n'))
            height=max(height,min(400,lines*17+6))
        ws.row_dimensions[row[0].row].height=height
    ws.auto_filter.ref=ws.dimensions;return ws


def workbook(settings,rows=None,report=None):
    wb=Workbook();wb.remove(wb.active)
    headers=layout(settings)['headers']
    types=layout(settings).get('column_types',{})
    ws=sheet(wb,'취합결과' if settings.get('template_snapshot') else '참석희망자',headers,rows or [],types)
    if rows is None:
        dv=DataValidation(type='list',formula1='"희망,미희망"');ws.add_data_validation(dv);dv.add('F2:F2001')
        sheet(wb,'작성안내',['항목','내용'],[['교육명',settings['course']],['교육일',settings['event_date']],['회신 마감일',settings['deadline']],['작성 방법','참석희망자 시트 1행 제목·순서를 유지하고 2행부터 입력'],['필수값','부서·성명·교육명·교육일·참석 희망(희망/미희망)'],['직급','선택 입력'],['희망자 없음','빈 파일 대신 담당자에게 없음 회신'],['개인정보','이름 외 연락처·주민번호 등을 추가하지 마세요.']])
    if report:
        # The tidy working table stays intact; submission and totals are derived sheets.
        submitted=[r for r in report['rows'] if r['decision']=='include' and attendee(r['values'],settings)]
        sheet(wb,'제출용 명단',['번호']+headers,[[i]+r['values'] for i,r in enumerate(submitted,1)],{str(int(k)+1):v for k,v in types.items()})
        reply_deps={f['id']:f['department'] for f in report['files']}
        sheet(wb,'부서별 집계',['부서','제출 대상 건수','회신 상태'],[[d['department'],sum((role_value(r['values'],settings,'department') or reply_deps.get(r['reply_id']))==d['department'] for r in submitted),d['status']] for d in report['departments']])
        sheet(wb,'미회신 목록',['부서','상태'],[[d['department'],d['status']] for d in report['departments'] if d['status']=='미회신'])
        has_merges=any(r.get('merge_provenance') for r in report['rows'])
        sheet(wb,'검토이력',headers+['처리','처리 사유','확인사항','원본 파일','원본 행']+(['병합 원본'] if has_merges else []),
          [r['values']+[r['decision'],r['reason'],' / '.join(r['issues']),r['filename'],(f"{r['group_provenance']['start']}~{r['group_provenance']['end']}" if r.get('group_provenance') else r['source_row'])]+
           (['; '.join(f"{p['header']}: {p['range']} ← {p['anchor']} ({p['value']})" for p in r.get('merge_provenance',[]))] if has_merges else []) for r in report['rows']],types)
        sheet(wb,'회신현황',['부서','상태','수신 행','포함 행','없음 회신 사유'],[[s['department'],s['status'],s['rows'],s['included'],(s['empty_reply'] or {}).get('reason','')] for s in report['departments']])
        sheet(wb,'출처',['원본 파일','회신 부서','SHA-256'],[[f['filename'],f['department'],f['sha256']] for f in report['files']])
        generic=layout(settings)['kind']=='generic'
        sheet(wb,'업무정보',['항목','내용'],[['업무',settings['title']],['자료명' if generic else '교육명',settings['course']],['기준일' if generic else '교육일',settings['event_date']],['마감일',settings['deadline']],['자료 버전',report['version']],['구분','작성 중 관리자료 · 최종 제출본 아님 · 제출용 시트에는 포함 검토한 대상만 표시' if report.get('preview') else '담당자 검토 후 저장한 자료 취합본' if generic else '담당자 검토 후 저장한 취합본 · 참석 확정 아님']])
        if settings.get('template_snapshot'):
            t=settings['template_snapshot']
            fields=('institution','name','version_label','sha256','sheet_name','header_row','header_end','header_merges','column_spans','fill_merged_columns','record_rows','group_columns','data_start','data_end','headers','roles','column_types','required','duplicate_keys')
            sheet(wb,'등록서식',['항목','등록 값'],[[k,encode(t.get(k)) if isinstance(t.get(k),(dict,list)) else t.get(k)] for k in fields])
            if t.get('record_rows',1)>1:
                sheet(wb,'행묶음 원본',['원본 파일','원본 행 범위','항목','읽기 기준','원본 셀·값','최초 취합값','검토 후 값'],
                  [[r['filename'],f"{g['start']}~{g['end']}",f['header'],('공통값 확인' if f['mode']=='consistent' else '여러 값 합치기' if f['mode']=='join' else f"묶음의 {int(f['mode'].split(':')[1])+1}번째 행만 읽기"),
                    '\n'.join(f"{c['cell']}: {c['value']}" for c in f['cells']),f['value'],r['values'][f['column']]]
                   for r in report['rows'] if (g:=r.get('group_provenance')) for f in g['fields']])
    stream=io.BytesIO();wb.save(stream);wb.close();return stream.getvalue()


def xlsx(raw,name):
    return Response(raw,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="{name}.xlsx"'})


def get_template(c,key):
    row=c.execute('SELECT * FROM collection_templates WHERE id=?',(key,)).fetchone()
    if not row:raise HTTPException(404,'등록한 취합 서식이 없습니다.')
    return row


def template_view(row):
    return {k:row[k] for k in ('id','institution','name','version_label','filename','sha256','created_at')}|{'config':json.loads(row['config'])}


async def template_file(file):
    try:raw=await file.read(MAX_BYTES+1)
    finally:await file.close()
    if len(raw)>MAX_BYTES or not (file.filename or '').lower().endswith('.xlsx'):raise HTTPException(422,'2MB 이하 .xlsx 파일만 지원합니다.')
    return raw


@router.post('/templates/inspect')
async def inspect_template(file:UploadFile=File(...)):
    raw=await template_file(file)
    try:return templates.inspect(raw)
    except ValueError as exc:raise HTTPException(422,str(exc)) from exc


@router.post('/templates/header-preview')
async def preview_header(file:UploadFile=File(...),sheet_name:str=Form(...),header_row:int=Form(...),header_end:int=Form(...)):
    raw=await template_file(file)
    try:return templates.preview(raw,sheet_name,header_row,header_end)
    except ValueError as exc:raise HTTPException(422,str(exc)) from exc


@router.get('/templates')
def list_templates():
    with db() as c:return [template_view(r) for r in c.execute('SELECT * FROM collection_templates ORDER BY created_at DESC,id')]


@router.post('/templates',status_code=201)
async def register_template(file:UploadFile=File(...),metadata:str=Form(...)):
    raw=await template_file(file)
    try:
        config=templates.Registration.model_validate_json(metadata)
        spec=templates.describe(raw,config)
    except ValueError as exc:raise HTTPException(422,str(exc)) from exc
    digest=hashlib.sha256(raw).hexdigest()
    with db() as c:
        c.execute('BEGIN IMMEDIATE')
        old=c.execute('SELECT * FROM collection_templates WHERE institution=? AND name=? AND version_label=?',(config.institution,config.name,config.version_label)).fetchone()
        if old:
            if old['sha256']==digest and json.loads(old['config'])==spec:return template_view(old)
            raise HTTPException(409,'같은 기관·서식명·버전이 있습니다. 변경한 양식은 새 버전으로 등록하세요.')
        key=uuid.uuid4().hex;filename=(file.filename or 'template.xlsx').replace('\\','/').split('/')[-1][:150]
        c.execute('INSERT INTO collection_templates VALUES (?,?,?,?,?,?,?,?,?)',(key,config.institution,config.name,config.version_label,encode(spec),filename,digest,raw,stamp()))
        return template_view(get_template(c,key))


@router.get('/templates/{key}/original')
def template_original(key:str):
    with db() as c:return xlsx(get_template(c,key)['original'],'registered-template-original')


@router.get('/jobs')
def jobs():
    with db() as c:return [dict(r)|{'settings':json.loads(r['settings'])} for r in c.execute('SELECT * FROM jobs ORDER BY created_at DESC')]


@router.post('/jobs',status_code=201)
def create(body:Create):
    key=uuid.uuid4().hex
    with db() as c:
        settings=body.model_dump(mode='json')
        if body.template_id:
            t=get_template(c,body.template_id)
            settings['template_snapshot']=json.loads(t['config'])|{'id':t['id'],'institution':t['institution'],'name':t['name'],'version_label':t['version_label'],'sha256':t['sha256']}
            if settings['template_snapshot']['kind']=='education' and body.deadline>body.event_date:raise HTTPException(422,'회신 마감일은 교육일 이후일 수 없습니다.')
        c.execute('INSERT INTO jobs VALUES (?,1,?,?)',(key,encode(settings),stamp()));log(c,key,'업무 생성',{'title':body.title})
        if layout(settings)['kind']=='education':
            c.execute('INSERT INTO job_workflows VALUES(?,?,?)',(key,0,encode(workflow.default())))
        return view(c,key)


@router.get('/jobs/{key}')
def read(key:str):
    with db() as c:return view(c,key)


@router.get('/jobs/{key}/related-folders')
def related_folders(key:str):
    with db() as c:job(c,key)
    return adhoc_workspace.related_folders(key)


@router.get('/jobs/{key}/workflow/references')
def workflow_references(key:str,q:str=''):
    with db() as c:
        job(c,key)
        return workflow.reference_options(c,q)


@router.patch('/jobs/{key}/workflow')
def workflow_update(key:str,body:workflow.Command):
    with db() as c:
        c.execute('BEGIN IMMEDIATE')
        v=view(c,key)
        workflow.update(c,v,body)
        return view(c,key)


@router.get('/jobs/{key}/template.xlsx')
def template(key:str):
    with db() as c:
        settings=job(c,key)['settings']
        if settings.get('template_snapshot'):return xlsx(get_template(c,settings['template_snapshot']['id'])['original'],'registered-template-original')
        return xlsx(workbook(settings),'education-template')


@router.get('/jobs/{key}/demo.zip')
def demo(key:str):
    with db() as c:settings=job(c,key)['settings']
    if settings.get('template_snapshot'):raise HTTPException(422,'합성 예제 ZIP은 기본 교육 서식용입니다. 등록 서식은 등록 원본으로 회신 파일을 준비하세요.')
    stream=io.BytesIO()
    with zipfile.ZipFile(stream,'w',zipfile.ZIP_DEFLATED) as z:
        for i,dep in enumerate(settings['departments']):
            row=[dep,f'모의직원{i+1}','주무관',settings['course'],settings['event_date'],'희망']
            rows=[row]
            if i==0:rows.append([dep,'모의추가직원','주무관',*row[3:]])
            if i==1:rows.extend([row[:],[dep,'','주무관',*row[3:]]])
            z.writestr(f'{i+1:02d}_합성회신.xlsx',workbook(settings,rows))
        bad=Workbook();bad.active.title='다른서식';bad.active.append(['이름','소속']);s=io.BytesIO();bad.save(s);bad.close()
        z.writestr('99_다른서식_거절예시.xlsx',s.getvalue())
        z.writestr('사용방법.txt','합성 가상 인물 자료입니다. 파일 번호 순서로 요청 대상 부서를 지정하세요.\n두 번째 부서는 중복 2행과 이름 누락 1행을 검토해야 합니다.\n마지막 파일은 다른 서식으로 거절되는 예시입니다. 실제 청구·선정·보고 자료로 사용하지 마세요.')
    return Response(stream.getvalue(),media_type='application/zip',headers={'Content-Disposition':'attachment; filename="synthetic-education-replies.zip"'})


@router.post('/jobs/{key}/files',status_code=201)
async def upload(key:str,version:int=Form(...),department:str=Form(''),file:UploadFile=File(...)):
    try:raw=await file.read(MAX_BYTES+1)
    finally:await file.close()
    if len(raw)>MAX_BYTES or not (file.filename or '').lower().endswith('.xlsx'):raise HTTPException(422,'2MB 이하 .xlsx 파일만 지원합니다.')
    digest=hashlib.sha256(raw).hexdigest()
    with db() as c:
        c.execute('BEGIN IMMEDIATE');j=job(c,key);settings=j['settings']
        old=c.execute('SELECT department FROM replies WHERE job_id=? AND sha256=?',(key,digest)).fetchone()
        if old:
            if department and old['department']!=department:raise HTTPException(409,'동일 파일이 다른 부서 회신으로 등록되어 있습니다. 등록 원본의 회신 부서를 수정하세요.')
            return view(c,key)
        if c.execute('SELECT COUNT(*) FROM replies WHERE job_id=?',(key,)).fetchone()[0]>=50:raise HTTPException(422,'업무별 최대 50개 파일까지 지원합니다.')
        try:parsed=templates.parse(raw,settings['template_snapshot']) if settings.get('template_snapshot') else parse(raw)
        except ValueError as exc:raise HTTPException(422,str(exc)) from exc
        parsed=[(r[0],r[1],r[2] if len(r)>2 else [],r[3] if len(r)>3 else None) for r in parsed]
        if not department:
            departments={role_value(values,settings,'department') for _,values,_,_ in parsed if role_value(values,settings,'department')}
            if len(departments)!=1 or not departments.issubset(set(settings['departments'])):
                raise HTTPException(422,'부서 열에서 회신 부서를 하나로 인식할 수 없습니다. 파일별 회신 부서를 지정하세요.')
            department=next(iter(departments))
        if department not in settings['departments']:raise HTTPException(422,'취합 대상 부서를 선택하세요.')
        if c.execute('SELECT COUNT(*) FROM rows WHERE job_id=?',(key,)).fetchone()[0]+len(parsed)>5000:raise HTTPException(422,'업무별 최대 5,000행까지 지원합니다.')
        bump(c,key,version);rid=uuid.uuid4().hex
        name=(file.filename or 'reply.xlsx').replace('\\','/').split('/')[-1][:150]
        c.execute('INSERT INTO replies VALUES (?,?,?,?,?,?,?)',(rid,key,department,name,digest,raw,stamp()))
        for source_row,values,provenance,group in parsed:
            row_id=uuid.uuid4().hex
            original_values=values[:]
            for item in provenance:
                if item['inherited'] and not group:original_values[item['column']]=''
            c.execute('INSERT INTO rows VALUES (?,?,?,?,?,?,?,?)',(row_id,key,rid,source_row,encode(original_values),encode(values),'pending',''))
            if provenance:c.execute('INSERT INTO row_merge_provenance VALUES (?,?)',(row_id,encode(provenance)))
            if group:c.execute('INSERT INTO row_group_provenance VALUES (?,?)',(row_id,encode(group)))
        log(c,key,'회신 파일 등록',{'file':name,'department':department,'rows':len(parsed),'sha256':digest})
        return view(c,key)


@router.patch('/jobs/{key}/rows/{row_id}')
def review(key:str,row_id:str,body:Review):
    with db() as c:
        c.execute('BEGIN IMMEDIATE');j=job(c,key)
        row=c.execute('SELECT r.*,f.department FROM rows r JOIN replies f ON f.id=r.reply_id WHERE r.id=? AND r.job_id=?',(row_id,key)).fetchone()
        if not row:raise HTTPException(404,'이 업무의 행을 찾을 수 없습니다.')
        values=[v.strip() for v in body.values]
        if len(values)!=len(layout(j['settings'])['headers']):raise HTTPException(422,'등록한 서식의 열 개수와 다릅니다.')
        errors=issues(values,j['settings'],row['department'])
        if body.decision=='include' and errors:raise HTTPException(422,'포함 전 수정 필요: '+' / '.join(errors))
        bump(c,key,body.version)
        c.execute('UPDATE rows SET values_json=?,decision=?,reason=? WHERE id=?',(encode(values),body.decision,body.reason,row_id))
        log(c,key,'행 검토',{'row_id':row_id,'decision':body.decision,'reason':body.reason,'before':json.loads(row['values_json']),'after':values})
        return view(c,key)


@router.patch('/jobs/{key}/files/{file_id}')
def change_reply_department(key:str,file_id:str,body:EmptyReply):
    with db() as c:
        c.execute('BEGIN IMMEDIATE');j=job(c,key)
        row=c.execute('SELECT department FROM replies WHERE id=? AND job_id=?',(file_id,key)).fetchone()
        if not row:raise HTTPException(404,'이 업무의 회신 파일이 없습니다.')
        if body.department not in j['settings']['departments']:raise HTTPException(422,'취합 대상 부서를 선택하세요.')
        if body.department==row['department']:raise HTTPException(422,'현재와 다른 회신 부서를 선택하세요.')
        bump(c,key,body.version)
        c.execute('UPDATE replies SET department=? WHERE id=?',(body.department,file_id))
        c.execute("UPDATE rows SET decision='pending',reason='' WHERE reply_id=?",(file_id,))
        log(c,key,'회신 부서 정정',{'file_id':file_id,'before':row['department'],'after':body.department,'reason':body.reason,'rows_reset':'pending'})
        return view(c,key)


@router.post('/jobs/{key}/empty-replies')
def empty(key:str,body:EmptyReply):
    with db() as c:
        c.execute('BEGIN IMMEDIATE');j=job(c,key)
        if body.department not in j['settings']['departments']:raise HTTPException(422,'취합 대상 부서를 선택하세요.')
        bump(c,key,body.version)
        c.execute('INSERT INTO empty_replies VALUES (?,?,?,?) ON CONFLICT(job_id,department) DO UPDATE SET reason=excluded.reason,created_at=excluded.created_at',(key,body.department,body.reason,stamp()))
        log(c,key,'희망자 없음 회신',body.model_dump());return view(c,key)


@router.post('/jobs/{key}/review-clean')
def review_clean(key:str,body:BatchReview):
    with db() as c:
        c.execute('BEGIN IMMEDIATE');v=view(c,key)
        clean=[r['id'] for r in v['rows'] if r['decision']=='pending' and not r['issues']]
        if not clean:raise HTTPException(422,'일괄 포함할 정상 보류 행이 없습니다.')
        bump(c,key,body.version)
        c.executemany("UPDATE rows SET decision='include',reason=? WHERE id=?",[(body.reason,rid) for rid in clean])
        log(c,key,'정상 행 일괄 포함',{'rows':clean,'reason':body.reason});return view(c,key)


@router.post('/jobs/{key}/empty-replies/{department}/cancel')
def cancel_empty(key:str,department:str,body:Version):
    with db() as c:
        c.execute('BEGIN IMMEDIATE');job(c,key)
        if not c.execute('SELECT 1 FROM empty_replies WHERE job_id=? AND department=?',(key,department)).fetchone():raise HTTPException(404,'없음 회신이 없습니다.')
        bump(c,key,body.version)
        c.execute('DELETE FROM empty_replies WHERE job_id=? AND department=?',(key,department))
        log(c,key,'희망자 없음 회신 취소',{'department':department});return view(c,key)


@router.get('/jobs/{key}/files/{file_id}')
def original(key:str,file_id:str):
    with db() as c:
        job(c,key);f=c.execute('SELECT original FROM replies WHERE job_id=? AND id=?',(key,file_id)).fetchone()
        if not f:raise HTTPException(404,'회신 파일이 없습니다.')
        return xlsx(f['original'],'original-reply')


@router.post('/jobs/{key}/finalize',status_code=201)
def finalize(key:str,body:Version):
    with db() as c:
        c.execute('BEGIN IMMEDIATE');v=view(c,key)
        if v['version']!=body.version:raise HTTPException(409,'다른 변경이 있습니다. 다시 열어 주세요.')
        if not v['can_finalize']:raise HTTPException(422,'미회신·보류 행·회신 충돌을 먼저 확인하세요.')
        if any(e['is_current'] for e in v['exports']):return v
        raw=workbook(v['settings'],[r['values'] for r in v['rows'] if r['decision']=='include' and attendee(r['values'],v['settings'])],v)
        eid=uuid.uuid4().hex
        c.execute('INSERT INTO exports VALUES (?,?,?,?,?,?)',(eid,key,v['version'],raw,hashlib.sha256(raw).hexdigest(),stamp()))
        log(c,key,'취합본 저장',{'export_id':eid,'version':v['version']});return view(c,key)


@router.get('/jobs/{key}/exports/{export_id}')
def download(key:str,export_id:str):
    with db() as c:
        job(c,key);r=c.execute('SELECT original FROM exports WHERE job_id=? AND id=?',(key,export_id)).fetchone()
        if not r:raise HTTPException(404,'저장한 취합본이 없습니다.')
        return xlsx(r['original'],'education-consolidated')

@router.get('/jobs/{key}/working.xlsx')
def working_report(key:str,version:int):
    with db() as c:
        v=view(c,key)
        if v['version']!=version:raise HTTPException(409,'취합 자료가 변경됐습니다. 화면을 새로 열어 주세요.')
        v['preview']=True
        rows=[r['values'] for r in v['rows'] if r['decision']=='include' and attendee(r['values'],v['settings'])]
        return xlsx(workbook(v['settings'],rows,v),'working-report-not-final')
