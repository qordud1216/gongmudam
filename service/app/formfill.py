"""서식 초안 채우기.

서식 정의(fields)는 다음 형식이다.
    {"key": "출장기간", "label": "출장 기간", "source": "base:출장 종료일" | "profile:성명" | "today" | "ask" | "",
     "hint": "예: 2026. 10. 13.~10. 15."}
값을 채우는 순서
1) profile: 설정 화면의 '내 정보'(성명·소속·직급 등)
2) base:   상황 안내에서 인식한 기준일
3) today:  작성일
4) LLM 연결 시: 담당자가 말한 상황 문장에서 나머지 칸을 뽑는다(상황 문장에 그대로 있는 글자만 받고, 아니면 비운다).
채운 결과는 화면에서 고친 뒤 DOCX로 내려받는다. 등록된 원본 서식 파일은 수정하지 않는다.
"""
from __future__ import annotations

import re
import unicodedata
from datetime import date, datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

from . import assistant, config, deadlines, llm, store

ROOT = Path(__file__).resolve().parent.parent
from .runtime_paths import state_root

OUTPUT = state_root() / "output"


def prefill(form: dict, *, query: str = "", base_values: dict[str, str] | None = None,
            today: date | None = None, use_llm: bool = True) -> dict:
    today = today or date.today()
    profile = store.get_setting("profile", {}) or {}
    bases = {k: deadlines.parse_iso(v) for k, v in (base_values or {}).items()}
    values: dict[str, str] = {}
    origin: dict[str, str] = {}
    for f in form["fields"]:
        src = f.get("source") or ""
        v = ""
        if src.startswith("profile:"):
            v = str(profile.get(src[8:], "") or "")
            if v:
                origin[f["key"]] = "내 정보"
        elif src.startswith("base:"):
            names = src[5:].split("~")
            ds = [bases.get(n) for n in names]
            if all(ds):
                v = "~".join(deadlines.fmt(d)[:-3] for d in ds)
                origin[f["key"]] = "상황에서 인식한 날짜"
        elif src == "today":
            v = deadlines.fmt(today)[:-3]
            origin[f["key"]] = "작성일"
        values[f["key"]] = v

    notice = None
    model = None
    empty = [f for f in form["fields"] if not values.get(f["key"]) and f.get("source") == "ask"]
    prof = config.active_profile()
    provider = llm.get_provider(prof)
    if use_llm and query and empty and provider.available:
        prompt = "\n".join([
            f"서식: {form['name']}", f"담당자가 말한 상황: {query}", "",
            "아래 칸 중 상황 문장에 나온 내용은 그 낱말을 그대로 옮겨 채우세요. 문장에 없는 값은 빈 문자열로 둡니다. 금액은 문장에 적힌 숫자를 그대로 옮기고 단위를 바꿔 계산하지 않습니다.",
            "예: 문장이 '부산 출장 다녀왔어요'이면 출장지 칸에 '부산'을 씁니다. '안전교육 때문에'면 목적 칸에 '안전교육'을 씁니다.",
            *[f"- {f['key']}: {f.get('hint', '')}" for f in empty],
            "", 'JSON 객체 하나로만 답하세요. 예: {"칸 이름": "값"}',
        ])
        try:
            raw = provider.chat([{"role": "system", "content": "당신은 공문서 서식 작성을 돕습니다. 지어내지 말고 주어진 문장에서만 값을 옮깁니다."},
                                 {"role": "user", "content": prompt}], max_tokens=500)
            data = llm.extract_json(raw)
            if isinstance(data, dict):
                for f in empty:
                    v = str(data.get(f["key"]) or "").strip()
                    if v and _grounded(v, query, f["key"]):
                        values[f["key"]] = v
                        origin[f["key"]] = "모델이 상황 문장에서 추출"
            model = prof.get("model")
        except llm.LLMError as e:
            notice = f"모델 서버 응답 오류로 빈칸을 채우지 못했습니다(오류 종류: {e.kind}). 직접 입력해 주세요."
    return {"form": form, "values": values, "origin": origin, "notice": notice, "model": model}


# 칸별 최대 글자 수(공백 제외). 금액 칸은 숫자와 '원'만 받는다.
_AMOUNT_KEYS = {"운임", "숙박비", "일비·식비", "합계", "계약 금액"}
_MAX_LEN = {"출장 신청 문서번호": 30, "규격·수량": 30}
_DEFAULT_MAX_LEN = 20
_AMOUNT = re.compile(r"\d+원?")


def _compact(text: str) -> str:
    """비교용으로 공백을 지우고 숫자 사이 천 단위 쉼표('59,800')를 없앤다."""
    text = re.sub(r"\s+", "", unicodedata.normalize("NFKC", text))
    return re.sub(r"(?<=\d),(?=\d)", "", text)


def _grounded(value: str, query: str, key: str = "") -> bool:
    """모델이 낸 값이 상황 문장 원문에 그대로 들어 있는 부분 문자열일 때만 승인한다(공백·천 단위 쉼표만 무시).
    모델이 새로 만든 문장, 바꾼 표현, 단위를 환산한 금액('5만 원' → 50000)은 어느 칸에도 넣지 않는다(직접 입력).
    원문 일부라도 칸에 맞지 않는 값(기한·일정 낱말이 든 값, 너무 긴 값, 숫자가 아닌 금액)은 버린다."""
    v = _compact(value)
    if not v or len(v) > _MAX_LEN.get(key, _DEFAULT_MAX_LEN) or v not in _compact(query):
        return False
    if key in _AMOUNT_KEYS:
        return bool(_AMOUNT.fullmatch(v))
    return not assistant._SCHEDULE.search(v)


def export_docx(form: dict, values: dict[str, str], *, card_title: str = "") -> Path:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "맑은 고딕"
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    style.font.size = Pt(11)
    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run(form["name"])
    r.bold = True
    r.font.size = Pt(18)
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for f in form["fields"]:
        row = table.add_row().cells
        row[0].text = f.get("label") or f["key"]
        row[1].text = values.get(f["key"], "")
        for p in row[0].paragraphs:
            for run in p.runs:
                run.bold = True
    doc.add_paragraph()
    note = doc.add_paragraph()
    nr = note.add_run(
        "※ 서무비서가 만든 초안입니다. 제출 전에 원본 서식과 내용을 확인하세요."
        + (f" (절차: {card_title})" if card_title else ""))
    nr.font.size = Pt(9)
    safe = re.sub(r'[\\/:*?"<>|\s]+', "_", form["name"]).strip("_")
    path = OUTPUT / f"{safe}_초안_{datetime.now():%Y%m%d_%H%M%S}.docx"
    doc.save(path)
    return path
