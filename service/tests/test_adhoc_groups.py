import io,unittest
from openpyxl import load_workbook
import test_adhoc_templates as base
from app import adhoc_templates


def grouped_source(names=('모의묶음1','모의묶음2'),department='총무팀',variant=None):
    rows=[]
    for i,name in enumerate(names,1):
        rows.extend([[f'{i:03}',name,department,'희망','AI 교육','2026-11-06','첫날 실습'],
                     ['',name,department,'희망','AI 교육','2026-11-06','둘째날 토론']])
    raw=base.source(rows,footer=True)
    if variant:
        wb=load_workbook(io.BytesIO(raw));variant(wb['기관회신']);s=io.BytesIO();wb.save(s);wb.close();raw=s.getvalue()
    return raw


class GroupedTemplateTest(unittest.TestCase):
    setUp=base.TemplateTest.setUp
    tearDown=base.TemplateTest.tearDown
    req=base.TemplateTest.req
    create=base.TemplateTest.create
    upload=base.TemplateTest.upload
    clean=base.TemplateTest.clean
    empty=base.TemplateTest.empty
    decide=base.TemplateTest.decide
    register=base.TemplateTest.register
    def metadata(self,**changes):
        return base.TemplateTest.metadata(self,record_rows=2,group_columns={6:'join'},version_label='3.0')|changes
    def setup_job(self,**changes):
        return self.create(template_id=self.register(base.source(),**changes)['id'])
    def test_two_physical_rows_one_record_and_exact_cells(self):
        j=self.upload(self.setup_job(),[],dep='',raw=grouped_source())
        self.assertEqual(j['summary']['total'],2)
        r=j['rows'][0];self.assertEqual(r['values'][6],'첫날 실습\n둘째날 토론')
        self.assertEqual((r['group_provenance']['start'],r['group_provenance']['end']),(5,6))
        self.assertEqual(r['group_provenance']['fields'][6]['cells'],[{'cell':'H5','value':'첫날 실습'},{'cell':'H6','value':'둘째날 토론'}])
        self.assertEqual(j['rows'][1]['source_row'],7)
    def test_conflicting_person_rejects_entire_file_atomically(self):
        j=self.setup_job();raw=grouped_source(variant=lambda ws:setattr(ws['C8'],'value','다른사람'))
        self.upload(j,[],raw=raw,status=422)
        after=self.req('GET',f"/jobs/{j['id']}");self.assertEqual(after['version'],1);self.assertFalse(after['files']);self.assertFalse(after['rows'])
    def test_explicit_row_offset_preserves_unselected_cell_evidence(self):
        j=self.upload(self.setup_job(group_columns={6:'row:1'}),[],raw=grouped_source())
        self.assertEqual(j['rows'][0]['values'][6],'둘째날 토론')
        self.assertEqual(len(j['rows'][0]['group_provenance']['fields'][6]['cells']),2)
    def test_blank_first_row_does_not_move_group_boundary(self):
        def change(ws):
            for c in ws[5]:c.value=None
        j=self.upload(self.setup_job(),[],raw=grouped_source(variant=change))
        self.assertEqual([r['source_row'] for r in j['rows']],[5,7]);self.assertEqual(j['rows'][0]['values'][1],'모의묶음1')
    def test_empty_blocks_are_skipped_without_shift(self):
        def change(ws):
            for row in ws.iter_rows(min_row=5,max_row=6):
                for c in row:c.value=None
        j=self.upload(self.setup_job(),[],raw=grouped_source(variant=change))
        self.assertEqual(len(j['rows']),1);self.assertEqual(j['rows'][0]['source_row'],7)
    def test_invalid_group_registration(self):
        for changes in [dict(record_rows=0),dict(record_rows=21),dict(data_end=None),dict(data_end=11),dict(group_columns={6:'row:2'}),dict(group_columns={6:'unknown'}),dict(group_columns={99:'join'}),dict(record_rows=1)]:
            self.register(status=422,**changes)
    def test_join_restricted_for_typed_role_and_duplicate_fields(self):
        for changes in [dict(group_columns={1:'join'}),dict(group_columns={6:'join'},column_types={6:'date'}),dict(group_columns={6:'join'},duplicate_keys=[6])]:self.register(status=422,**changes)
    def test_three_row_records_and_registered_boundary(self):
        raw=base.source([['001','모의셋','총무팀','','','','첫째'],['','','','','','','둘째'],['','','','','','','셋째']])
        j=self.upload(self.setup_job(record_rows=3,data_end=10),[],raw=raw)
        self.assertEqual(len(j['rows']),1);self.assertEqual(j['rows'][0]['values'][6],'첫째\n둘째\n셋째')
    def test_merge_must_stay_inside_one_group(self):
        def change(ws):ws.merge_cells('D6:D7')
        j=self.setup_job(fill_merged_columns=[2]);self.upload(j,[],raw=grouped_source(variant=change),status=422)
    def test_merged_value_on_selected_blank_row_is_resolved(self):
        def change(ws):ws.merge_cells('D5:D6')
        j=self.upload(self.setup_job(fill_merged_columns=[2],group_columns={2:'row:1',6:'join'}),[],raw=grouped_source(variant=change))
        self.assertEqual(j['rows'][0]['values'][2],'총무팀');self.assertTrue(j['rows'][0]['merge_provenance'])
    def test_formula_in_unselected_row_is_rejected(self):
        j=self.setup_job(group_columns={6:'row:0'});self.upload(j,[],raw=grouped_source(variant=lambda ws:setattr(ws['H6'],'value','=1+1')),status=422)
    def test_blank_reply_is_not_a_record(self):
        j=self.setup_job();self.upload(j,[],raw=base.source(),status=422)
    def test_duplicate_keys_compare_groups(self):
        j=self.upload(self.setup_job(),[],raw=grouped_source(names=('同名','同名')))
        self.assertTrue(all(any('중복' in i for i in r['issues']) for r in j['rows']))
    def test_review_export_original_cells_and_previous_export_immutable(self):
        j=self.upload(self.setup_job(),[],raw=grouped_source());j=self.clean(j);j=self.empty(j,'행정팀');j=self.empty(j,'안전팀')
        saved=self.req('POST',f"/jobs/{j['id']}/finalize",dict(version=j['version']),201)
        url=f"/api/adhoc/jobs/{j['id']}/exports/{saved['exports'][0]['id']}";before=self.client.get(url).content
        wb=load_workbook(io.BytesIO(before));self.assertEqual(wb['취합결과'].max_row,3);self.assertEqual(wb['취합결과']['G2'].value,'첫날 실습\n둘째날 토론')
        self.assertEqual(wb['검토이력']['L2'].value,'5~6');self.assertEqual(wb['행묶음 원본']['B2'].value,'5~6')
        self.assertIn('H6: 둘째날 토론',wb['행묶음 원본']['E8'].value);wb.close()
        values=j['rows'][0]['values'][:];values[6]='담당자 확인 수정';updated=self.decide(j,j['rows'][0],'include',values)
        self.assertEqual(updated['rows'][0]['group_provenance']['fields'][6]['value'],'첫날 실습\n둘째날 토론');self.assertEqual(self.client.get(url).content,before)
    def test_new_template_version_does_not_change_existing_group_rules(self):
        j=self.setup_job();self.register(version_label='3.1',group_columns={6:'row:0'})
        v=self.req('GET',f"/jobs/{j['id']}");self.assertEqual(v['layout']['group_columns'],{'6':'join'})

if __name__=='__main__':unittest.main()
