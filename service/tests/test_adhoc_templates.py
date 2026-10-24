import io,json,unittest,zipfile
from openpyxl import Workbook,load_workbook
import test_adhoc
from app import adhoc


HEADERS=['순번','이름','소속부서','희망여부','교육과정','교육일자','비고']


def source(rows=(),headers=None,header=4,footer=False,extra_sheet=True):
    wb=Workbook();ws=wb.active;ws.title='기관회신';ws.merge_cells('A1:H1');ws['A1']='[합성] 기관에서 기존에 쓰던 교육 회신서';ws['A2']='이름·소속을 확인하여 회신'
    for i,v in enumerate(headers or HEADERS,2):ws.cell(header,i,v)
    for n,row in enumerate(rows,header+1):
        for i,v in enumerate(row,2):ws.cell(n,i,v)
    if footer:ws.cell(15,2,'합계·확인 서명란')
    if extra_sheet:wb.create_sheet('기관안내')['A1']='선택하지 않은 안내 시트'
    stream=io.BytesIO();wb.save(stream);wb.close();return stream.getvalue()


class TemplateTest(unittest.TestCase):
    setUp=test_adhoc.AdhocTest.setUp
    tearDown=test_adhoc.AdhocTest.tearDown
    req=test_adhoc.AdhocTest.req
    create=test_adhoc.AdhocTest.create
    upload=test_adhoc.AdhocTest.upload
    clean=test_adhoc.AdhocTest.clean
    empty=test_adhoc.AdhocTest.empty
    decide=test_adhoc.AdhocTest.decide

    def metadata(self,**changes):
        return dict(institution='[합성] 한빛기관',name='기관 기존 회신서',version_label='1.0',sheet_name='기관회신',header_row=4,data_start=5,data_end=12,kind='generic',roles={'department':2,'name':1},required=[1,2],duplicate_keys=[1,2],selected_area_confirmed=True)|changes
    def register(self,raw=None,status=201,**changes):
        r=self.client.post('/api/adhoc/templates',data={'metadata':json.dumps(self.metadata(**changes),ensure_ascii=False)},files={'file':('기관서식.xlsx',raw if raw is not None else source(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
        self.assertEqual(r.status_code,status,r.text);return r.json()
    def row(self,name='가상하나',dep='총무팀',wish='참석'):return ['001',name,dep,wish,'AI 교육','2026-11-06','기관 전용 비고 보존']

    def test_inspect_and_register_original_and_headers(self):
        raw=source();r=self.client.post('/api/adhoc/templates/inspect',files={'file':('서식.xlsx',raw)})
        self.assertEqual(r.status_code,200);self.assertEqual(r.json()['sheets'][0]['preview'][3][1:],HEADERS)
        t=self.register(raw);self.assertEqual(t['config']['start_column'],2);self.assertEqual(t['config']['headers'],HEADERS)
        self.assertEqual(t['config']['ignored_sheets'],['기관안내'])
        self.assertEqual(self.client.get(f"/api/adhoc/templates/{t['id']}/original").content,raw)

    def test_version_and_identical_retry(self):
        raw=source();t=self.register(raw);again=self.register(raw);self.assertEqual(t,again)
        self.register(raw,status=409,required=[0,1,2])
        newer=self.register(raw,version_label='2.0');self.assertNotEqual(t['id'],newer['id'])
        adhoc.initialize();self.assertEqual(len(self.req('GET','/templates')),2)

    def test_invalid_header_mapping_and_range(self):
        for changes in [dict(header_row=1),dict(roles={'name':1,'department':1}),dict(required=[99]),dict(data_start=4),dict(data_end=3),dict(selected_area_confirmed=False)]:self.register(status=422,**changes)
        self.register(source(headers=['A','A','B','C','D','E','F']),status=422)
        self.register(kind='education',status=422)

    def test_generic_merge_all_columns_extra_sheet_footer_and_snapshot(self):
        raw=source(footer=True);t=self.register(raw);j=self.create(template_id=t['id'],event_date='2026-01-01')
        self.assertEqual(j['layout']['headers'],HEADERS)
        self.assertEqual(self.client.get(f"/api/adhoc/jobs/{j['id']}/template.xlsx").content,raw)
        j=self.upload(j,[],raw=source([self.row()],footer=True),dep='')
        self.assertEqual(j['rows'][0]['source_row'],5);self.assertEqual(j['rows'][0]['values'],self.row())
        j=self.clean(j);j=self.empty(self.empty(j,'행정팀'),'안전팀')
        final=self.req('POST',f"/jobs/{j['id']}/finalize",dict(version=j['version']),201)
        r=self.client.get(f"/api/adhoc/jobs/{j['id']}/exports/{final['exports'][0]['id']}")
        wb=load_workbook(io.BytesIO(r.content));self.assertEqual([c.value for c in wb['취합결과'][1]],HEADERS)
        self.assertEqual(wb['취합결과'].cell(2,7).value,'기관 전용 비고 보존');self.assertEqual(wb['취합결과'].cell(2,1).value,'001')
        self.assertEqual(wb['취합결과'].cell(2,1).number_format,'@')
        self.assertIn('등록서식',wb.sheetnames);wb.close()
        self.register(version_label='2.0',duplicate_keys=[0])
        frozen=self.req('GET',f"/jobs/{j['id']}")['settings']['template_snapshot'];self.assertEqual(frozen['version_label'],'1.0');self.assertEqual(frozen['duplicate_keys'],[1,2])

    def test_education_custom_roles_and_tokens(self):
        t=self.register(kind='education',roles={'department':2,'name':1,'course':4,'date':5,'wish':3},positive='참석',negative='불참')
        j=self.create(template_id=t['id']);j=self.upload(j,[],dep='',raw=source([self.row(),self.row(name='가상둘',wish='불참')]))
        self.assertFalse(j['rows'][0]['issues']);j=self.clean(j);self.assertEqual(j['summary']['attendees'],1)
        self.assertEqual(j['summary']['included'],2)

    def test_unmatched_shape_and_formula_atomic(self):
        t=self.register();j=self.create(template_id=t['id'])
        self.upload(j,[],raw=source([self.row()],headers=list(reversed(HEADERS))),status=422)
        self.upload(j,[],raw=source([self.row()],header=3),status=422)
        row=self.row();row[6]='=1+1';self.upload(j,[],raw=source([row]),status=422)
        self.assertEqual(self.req('GET',f"/jobs/{j['id']}")['version'],1)

    def test_registered_required_and_duplicate_keys(self):
        t=self.register();j=self.create(template_id=t['id']);rows=[self.row(),self.row(),self.row(name='')]
        j=self.upload(j,[],raw=source(rows),dep='');self.assertEqual(j['summary']['clean_pending'],0)
        self.assertTrue(any('중복 의심' in x for x in j['rows'][0]['issues']));self.assertIn('필수값 누락',j['rows'][2]['issues'])

    def test_missing_department_role_manual_assignment(self):
        t=self.register(roles={},required=[1]);j=self.create(template_id=t['id'])
        self.upload(j,[],raw=source([self.row()]),dep='',status=422)
        j=self.upload(j,[],raw=source([self.row()]),dep='행정팀');self.assertFalse(j['rows'][0]['issues'])

    def test_row_count_contract_and_job_template_not_found(self):
        self.req('POST','/jobs',self.settings|{'template_id':'missing'},404)
        t=self.register();j=self.upload(self.create(template_id=t['id']),[],raw=source([self.row()]))
        self.req('PATCH',f"/jobs/{j['id']}/rows/{j['rows'][0]['id']}",dict(version=j['version'],values=['only one'],decision='include',reason='열 개수를 확인한 사유'),422)

    def test_extra_data_columns_and_merged_body_rejected(self):
        t=self.register();j=self.create(template_id=t['id']);raw=source([self.row()]);wb=load_workbook(io.BytesIO(raw));wb['기관회신']['I5']='새 열';s=io.BytesIO();wb.save(s);wb.close();self.upload(j,[],raw=s.getvalue(),status=422)
        wb=load_workbook(io.BytesIO(raw));wb['기관회신'].merge_cells('B5:C5');s=io.BytesIO();wb.save(s);wb.close();self.upload(j,[],raw=s.getvalue(),status=422)

    def test_number_and_date_types_preserved_and_invalid_blocked(self):
        t=self.register(column_types={0:'number',5:'date'})
        j=self.upload(self.create(template_id=t['id']),[],raw=source([self.row()]))
        j=self.clean(j);raw=adhoc.workbook(j['settings'],[j['rows'][0]['values']],j)
        wb=load_workbook(io.BytesIO(raw));self.assertEqual(wb['취합결과']['A2'].data_type,'n');self.assertTrue(wb['취합결과']['F2'].is_date);wb.close()
        bad=self.row(name='가상둘');bad[0]='번호없음';j=self.upload(j,[],raw=source([bad]))
        self.assertTrue(any('숫자' in s for s in j['rows'][-1]['issues']))

    def test_general_numeric_negative_is_valid_and_snapshot_contract_immutable(self):
        t=self.register(column_types={0:'number'});row=self.row();row[0]='-2.5'
        j=self.upload(self.create(template_id=t['id']),[],raw=source([row]))
        self.assertFalse(j['rows'][0]['issues'])
        self.register(version_label='2.0',column_types={0:'text'})
        self.assertEqual(self.req('GET',f"/jobs/{j['id']}")['layout']['column_types'],{'0':'number'})


if __name__=='__main__':unittest.main()
