import io,json,tempfile,unittest,hashlib,os
from pathlib import Path
from unittest.mock import patch
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook
from app import folders


def doc(title='봄 행사',year='2026',day='2026-04-10',extra=''):
 return f'행사명: {title}\n연도: {year}\n개최일: {day}\n주관부서: 총무팀\n{extra}'
class FolderTest(unittest.TestCase):
 def setUp(self):
  self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name);self.a=self.root/'자료';self.b=self.root/'다운로드';self.a.mkdir();self.b.mkdir()
  self.patch=patch.object(folders,'DB_PATH',self.root/'folders.db');self.patch.start();folders.initialize();app=FastAPI();app.include_router(folders.router);self.client=TestClient(app)
 def tearDown(self):self.client.close();self.patch.stop();self.temp.cleanup()
 def write(self,folder,name,text):p=folder/name;p.write_text(text,encoding='utf-8');return p
 def req(self,method,path,body=None,status=200):
  r=self.client.request(method,'/api/folders'+path,json=body);self.assertEqual(r.status_code,status,r.text);return r.json()
 def connect(self):return self.req('POST','/libraries',dict(name='[합성] 업무 자료함',roots=[str(self.a),str(self.b)]),201)
 def scan(self,v):return self.req('POST',f"/libraries/{v['id']}/scan",dict(version=v['version']))
 def test_scattered_and_mixed_folders_are_grouped_by_content(self):
  self.write(self.a,'계획.txt',doc());self.write(self.b,'회신.txt',doc());self.write(self.a,'가을_계획.txt',doc('가을 행사',day='2026-10-10'))
  v=self.connect();self.assertEqual(v['summary']['projects'],2);self.assertEqual(sorted(p['files'] for p in v['projects']),[1,2]);self.assertTrue(all(f['clues']['evidence'] for f in v['files']))
 def test_identical_title_different_years_stay_apart(self):
  self.write(self.a,'계획.txt',doc());self.write(self.b,'계획.txt',doc(year='2025',day='2025-04-10'))
  self.assertEqual(self.connect()['summary']['projects'],2)
 def test_identical_title_different_dates_stay_apart_and_ambiguous_unassigned(self):
  self.write(self.a,'a.txt',doc());self.write(self.a,'b.txt',doc(day='2026-04-11'));self.write(self.b,'회신.txt','행사명: 봄 행사\n연도: 2026')
  v=self.connect();self.assertEqual(v['summary']['projects'],2);f=next(f for f in v['files'] if f['name']=='회신.txt');self.assertFalse(f['project_ids']);self.assertIn('여러 개최일',f['clues']['reason'])
 def test_one_dated_candidate_can_receive_undated_reply(self):
  self.write(self.a,'계획.txt',doc());self.write(self.b,'회신.txt','행사명: 봄 행사\n연도: 2026')
  v=self.connect();self.assertEqual(v['summary']['projects'],1);self.assertEqual(v['projects'][0]['files'],2)
 def test_different_organizers_not_merged(self):
  self.write(self.a,'a.txt',doc());self.write(self.b,'b.txt',doc().replace('총무팀','행정팀'));self.assertEqual(self.connect()['summary']['projects'],2)
 def test_multiple_names_in_document_require_review(self):
  self.write(self.a,'계획.txt',doc(extra='행사명: 가을 행사'));v=self.connect();self.assertFalse(v['files'][0]['project_ids']);self.assertIn('섞임',v['files'][0]['clues']['reason'])
 def test_common_and_unrecognized_never_guessed_from_directory(self):
  self.write(self.a,'공통_서식.txt','표준 양식');self.write(self.b,'봄 행사 회신.txt','참석자 김모의')
  v=self.connect();self.assertEqual(v['summary']['common'],1);self.assertEqual(v['summary']['projects'],0);self.assertTrue(all(not f['confirmed'] for f in v['files']))
 def test_duplicate_content_across_folders_retained(self):
  self.write(self.a,'a.txt',doc());self.write(self.b,'b.txt',doc());v=self.connect();self.assertEqual(v['summary']['files'],2);self.assertTrue(all(len(f['duplicates'])==1 for f in v['files']))
 def test_overlapping_roots_do_not_duplicate_file(self):
  child=self.a/'하위';child.mkdir();self.write(child,'계획.txt',doc())
  v=self.req('POST','/libraries',dict(name='중복 경로',roots=[str(self.a),str(child)]),201);self.assertEqual(v['summary']['files'],1)
 def test_manual_shared_assignment_and_change_survive_rescan(self):
  source=self.write(self.a,'계획.txt',doc());self.write(self.b,'가을.txt',doc('가을 행사',day='2026-10-10'));v=self.connect();f=next(f for f in v['files'] if f['name']=='계획.txt');ids=[p['id'] for p in v['projects']]
  v=self.req('PATCH',f"/libraries/{v['id']}/files/{f['id']}",dict(version=v['version'],project_ids=ids,reason='두 행사에서 같이 사용 확인'))
  v=self.scan(v);f=next(x for x in v['files'] if x['id']==f['id']);self.assertTrue(f['confirmed']);self.assertEqual(f['project_ids'],ids)
  source.write_text(doc(extra='내용 변경'),encoding='utf-8');v=self.scan(v);f=next(x for x in v['files'] if x['id']==f['id']);self.assertTrue(f['confirmed']);self.assertTrue(f['content_changed']);self.assertTrue(f['changed']);self.assertEqual(f['project_ids'],ids)
 def test_missing_file_kept_as_unconfirmed_history(self):
  p=self.write(self.a,'계획.txt',doc());v=self.connect();p.rename(self.root/'outside.txt');v=self.scan(v);self.assertEqual(v['summary']['missing'],1);self.assertEqual(v['summary']['files'],0);self.assertFalse(v['files'][0]['present'])
 def test_add_folder_without_recreating_library(self):
  self.write(self.a,'a.txt',doc());v=self.req('POST','/libraries',dict(name='추가 연결',roots=[str(self.a)]),201);self.write(self.b,'b.txt',doc())
  v=self.req('POST',f"/libraries/{v['id']}/scan",dict(version=v['version'],roots=[str(self.a),str(self.b)]));self.assertEqual(v['projects'][0]['files'],2)
 def test_stale_writes_atomic_and_unknown_project_rejected(self):
  self.write(self.a,'계획.txt',doc());v=self.connect();pid=v['projects'][0]['id'];path=f"/libraries/{v['id']}/projects/{pid}/confirm"
  result=self.req('POST',path,dict(version=v['version']));self.req('POST',path,dict(version=v['version']),409)
  self.req('PATCH',f"/libraries/{v['id']}/files/{v['files'][0]['id']}",dict(version=result['version'],project_ids=['missing'],reason='잘못된 업무 분류'),422)
  self.assertEqual(self.req('GET',f"/libraries/{v['id']}")['version'],result['version'])
 def test_project_name_confirm_and_manual_new_project(self):
  self.write(self.a,'계획.txt',doc());v=self.connect();pid=v['projects'][0]['id'];v=self.req('PATCH',f"/libraries/{v['id']}/projects/{pid}",dict(version=v['version'],name='봄 행사 사용자 이름'))
  v=self.req('POST',f"/libraries/{v['id']}/projects/{pid}/confirm",dict(version=v['version']));v=self.scan(v);self.assertEqual(v['projects'][0]['name'],'봄 행사 사용자 이름');self.assertTrue(v['files'][0]['confirmed'])
  v=self.req('POST',f"/libraries/{v['id']}/projects",dict(version=v['version'],name='직접 분류 업무'),201);self.assertEqual(len(v['projects']),2)
 def test_stage_file_is_evidence_not_completion(self):
  self.write(self.a,'결과보고.txt',doc());v=self.connect();p=v['projects'][0];self.assertEqual(p['stages'][3]['files'],[v['files'][0]['id']]);self.assertIn('미처리로 확정하지 않습니다',p['next_actions'][1])
 def test_file_types_errors_and_formula_exclusion(self):
  self.write(self.a,'알수없음.bin','행사명: 숨은 행사');self.write(self.a,'손상.docx','bad zip');self.write(self.a,'텍스트.txt',doc())
  wb=Workbook();ws=wb.active;ws['A1']='행사명: 엑셀 행사';ws['A2']='연도: 2026';ws['A3']='=HYPERLINK("http://invalid")';wb.save(self.a/'표.xlsx');wb.close()
  v=self.connect();self.assertEqual(v['summary']['files'],4);excel=next(f for f in v['files'] if f['name']=='표.xlsx');self.assertTrue(excel['project_ids']);self.assertNotIn('HYPERLINK',json.dumps(excel['units']))
  self.assertIn('실패',next(f for f in v['files'] if f['name']=='손상.docx')['status'])
 def test_limits_are_reported_and_roots_rejected(self):
  self.req('POST','/libraries',dict(name='상대 경로',roots=['relative']),422);self.req('POST','/libraries',dict(name='없는 경로',roots=[str(self.root/'없음')]),422)
  self.req('POST','/libraries',dict(name='드라이브',roots=[str(self.root.anchor)]),422)
  self.write(self.a,'a.txt',doc());self.write(self.a,'b.txt',doc())
  with patch.object(folders,'MAX_FILES',1):v=self.connect()
  self.assertEqual(v['summary']['files'],1);self.assertTrue(v['warnings'])
 def test_original_hash_and_mtime_unchanged_and_report_export(self):
  p=self.write(self.a,'계획.txt',doc());before=(p.read_bytes(),p.stat().st_mtime_ns);v=self.connect();self.scan(v)
  self.assertEqual(before,(p.read_bytes(),p.stat().st_mtime_ns));r=self.client.get(f"/api/folders/libraries/{v['id']}/report.json");self.assertEqual(r.status_code,200);self.assertEqual(r.json()['files'][0]['sha256'],hashlib.sha256(before[0]).hexdigest())
 def test_unsupported_file_without_body_stays_review(self):
  self.write(self.a,'2026 봄 행사.hwp','fake');v=self.connect();self.assertEqual(v['files'][0]['status'],'본문 미지원 · 파일명만 확인');self.assertFalse(v['files'][0]['project_ids'])
 def test_common_confirmation_survives_rescan(self):
  self.write(self.a,'공통서식.txt','빈 서식');v=self.connect();f=v['files'][0]
  v=self.req('PATCH',f"/libraries/{v['id']}/files/{f['id']}",dict(version=v['version'],bucket='common',project_ids=[],reason='부서 공용서식으로 확인'))
  v=self.scan(v);self.assertTrue(v['files'][0]['confirmed']);self.assertEqual(v['summary']['common'],1)

 def test_xlsx_separate_label_cells_and_table_headers(self):
  wb=Workbook();ws=wb.active;ws.append(['행사명','셀 행사']);ws.append(['연도','2026']);wb.save(self.a/'라벨.xlsx');wb.close()
  wb=Workbook();ws=wb.active;ws.append(['행사명','연도','개최일']);ws.append(['표 행사','2026','2026-10-30']);wb.save(self.b/'표.xlsx');wb.close()
  v=self.connect();self.assertEqual(v['summary']['projects'],2);self.assertTrue(all(f['project_ids'] for f in v['files']))
 def test_existing_collection_job_link_does_not_modify_job(self):
  from app import adhoc
  with patch.object(adhoc,'DB_PATH',self.root/'adhoc.db'):
   adhoc.initialize()
   with adhoc.db() as c:c.execute('INSERT INTO jobs VALUES(?,?,?,?)',('unit-job',1,json.dumps({'title':'합성 자료 취합'}),'unit'))
   self.write(self.a,'계획.txt',doc());v=self.connect();pid=v['projects'][0]['id'];path=f"/libraries/{v['id']}/projects/{pid}/links"
   v=self.req('POST',path,dict(version=v['version'],job_id='unit-job'));self.assertEqual(v['projects'][0]['links'][0]['title'],'합성 자료 취합')
   v=self.req('POST',path,dict(version=v['version'],job_id='unit-job'));self.assertEqual(len(v['projects'][0]['links']),1)
   self.req('POST',path,dict(version=v['version'],job_id='unknown'),404)
   with adhoc.db() as c:self.assertEqual(adhoc.job(c,'unit-job')['version'],1)
 def test_docx_snapshot_parsing_and_empty_pdf(self):
  import zipfile
  from pypdf import PdfWriter
  xml='<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>행사명: 문서 행사</w:t></w:r></w:p><w:p><w:r><w:t>연도: 2026</w:t></w:r></w:p></w:body></w:document>'
  with zipfile.ZipFile(self.a/'문서.docx','w') as z:z.writestr('word/document.xml',xml)
  writer=PdfWriter();writer.add_blank_page(width=100,height=100)
  with (self.b/'빈스캔.pdf').open('wb') as f:writer.write(f)
  v=self.connect();document=next(f for f in v['files'] if f['name']=='문서.docx');self.assertTrue(document['project_ids']);self.assertEqual(document['clues']['evidence'][0]['loc'],'1번째 문단')
  self.assertFalse(next(f for f in v['files'] if f['name']=='빈스캔.pdf')['project_ids'])

if __name__=='__main__':unittest.main()
