import io,unittest
from openpyxl import Workbook,load_workbook
import test_adhoc_templates
from app import adhoc_templates


def merged_source(department='총무팀',names=(),variant=None):
    wb=Workbook();ws=wb.active;ws.title='기관회신';ws.merge_cells('A1:H1');ws['A1']='[합성] 다단 제목·병합 부서 교육회신'
    for area,value in [('B4:B6','순번'),('C4:D4','인적사항'),('C5:C6','성명'),('D5:D6','소속부서'),
      ('E4:F4','교육정보'),('E5:E6','교육명'),('F5:F6','교육일'),('G4:G6','참석\n여부'),('H4:H6','비고')]:
        ws.merge_cells(area);ws.cell(ws[area.split(':')[0]].row,ws[area.split(':')[0]].column,value)
    for row,name in enumerate(names,7):
        values=[f'{row-6:03}',name,department,'공공 AI 실무교육','2026-11-06','희망',f'병합 부서 회신 {row-6}']
        for col,value in enumerate(values,2):ws.cell(row,col,value)
    ws['D7']=department if names else None;ws.merge_cells('D7:D8');ws['B15']='꼬리말 · 담당자 확인'
    wb.create_sheet('작성안내')['A1']='선택 시트 밖 안내'
    if variant:variant(ws)
    stream=io.BytesIO();wb.save(stream);wb.close();return stream.getvalue()


def wide_source(department='총무팀',names=(),variant=None):
    wb=load_workbook(io.BytesIO(merged_source(department,names)));ws=wb['기관회신']
    ws.unmerge_cells('H4:H6');ws.merge_cells('H4:I6')
    ws.merge_cells('H7:I7');ws.merge_cells('H8:I8')
    if variant:variant(ws)
    stream=io.BytesIO();wb.save(stream);wb.close();return stream.getvalue()


class MergedTemplateTest(unittest.TestCase):
    setUp=test_adhoc_templates.TemplateTest.setUp
    tearDown=test_adhoc_templates.TemplateTest.tearDown
    req=test_adhoc_templates.TemplateTest.req
    create=test_adhoc_templates.TemplateTest.create
    upload=test_adhoc_templates.TemplateTest.upload
    clean=test_adhoc_templates.TemplateTest.clean
    empty=test_adhoc_templates.TemplateTest.empty
    decide=test_adhoc_templates.TemplateTest.decide
    register=test_adhoc_templates.TemplateTest.register
    def metadata(self,**changes):
        return dict(institution='[합성] 한빛기관',name='병합 교육회신',version_label='2.0',sheet_name='기관회신',header_row=4,header_end=6,
          data_start=7,data_end=12,kind='education',roles=dict(name=1,department=2,course=3,date=4,wish=5),required=[1,2],
          duplicate_keys=[1,2,3,4],fill_merged_columns=[2],selected_area_confirmed=True)|changes
    def setup_job(self):
        t=self.register(merged_source());return self.create(template_id=t['id'],course='공공 AI 실무교육')

    def test_merged_hierarchy_and_multiline_cell(self):
        t=self.register(merged_source());cfg=t['config']
        self.assertEqual(cfg['headers'],['순번','인적사항 / 성명','인적사항 / 소속부서','교육정보 / 교육명','교육정보 / 교육일','참석 여부','비고'])
        self.assertEqual(cfg['header_paths'][1],['인적사항','성명']);self.assertEqual(cfg['header_end'],6)
        self.assertIn('C4:D4',cfg['header_merges']);self.assertEqual(cfg['start_column'],2)

    def test_preview_and_registered_contract_same(self):
        raw=merged_source();r=self.client.post('/api/adhoc/templates/header-preview',data=dict(sheet_name='기관회신',header_row=4,header_end=6),files={'file':('양식.xlsx',raw)})
        self.assertEqual(r.status_code,200);self.assertEqual(r.json()['headers'],self.register(raw)['config']['headers'])

    def test_vertical_body_fill_preserves_physical_original_and_provenance(self):
        j=self.upload(self.setup_job(),[],raw=merged_source(names=['가상하나','가상둘']),dep='')
        self.assertEqual(len(j['rows']),2);first,second=j['rows']
        self.assertEqual(second['values'][2],'총무팀');self.assertEqual(second['original_values'][2],'')
        p=second['merge_provenance'][0];self.assertEqual((p['range'],p['anchor'],p['inherited']),('D7:D8','D7',True))
        self.assertFalse(first['merge_provenance'][0]['inherited']);self.assertEqual(second['source_row'],8)
        self.assertEqual(j['summary']['clean_pending'],2)

    def test_two_files_consolidate_and_audit_anchor_exported(self):
        j=self.upload(self.setup_job(),[],raw=merged_source(names=['가상하나','가상둘']),dep='')
        j=self.upload(j,[],raw=merged_source('행정팀',['가상셋','가상넷']),dep='');j=self.clean(j);j=self.empty(j,'안전팀')
        final=self.req('POST',f"/jobs/{j['id']}/finalize",dict(version=j['version']),201)
        raw=self.client.get(f"/api/adhoc/jobs/{j['id']}/exports/{final['exports'][0]['id']}").content
        wb=load_workbook(io.BytesIO(raw));self.assertEqual(wb['취합결과'].max_row,5)
        self.assertEqual(wb['취합결과']['C3'].value,'총무팀')
        headers=[c.value for c in wb['검토이력'][1]];idx=headers.index('병합 원본')+1
        self.assertIn('D7:D8',wb['검토이력'].cell(3,idx).value);self.assertIn('D7',wb['검토이력'].cell(3,idx).value);wb.close()

    def test_wrong_header_merge_structure_rejected_even_same_labels(self):
        j=self.setup_job()
        def change(ws):ws.unmerge_cells('C5:C6');ws['C6']='성명'
        raw=merged_source(names=['가상하나','가상둘'],variant=change)
        self.upload(j,[],raw=raw,status=422)
        self.assertEqual(self.req('GET',f"/jobs/{j['id']}")['version'],1)

    def test_changed_header_text_rejected(self):
        j=self.setup_job();self.upload(j,[],raw=merged_source(names=['가상하나','가상둘'],variant=lambda ws:setattr(ws['C4'],'value','다른 그룹')),status=422)

    def test_vertical_fill_requires_explicit_registered_columns(self):
        self.register(merged_source(),fill_merged_columns=[],status=422)
        self.assertEqual(self.req('GET','/templates'),[])

    def test_horizontal_body_merge_rejected(self):
        self.register(merged_source(variant=lambda ws:ws.merge_cells('E7:F7')),status=422)
        j=self.setup_job();self.upload(j,[],raw=merged_source(names=['가상하나','가상둘'],variant=lambda ws:ws.merge_cells('E7:F7')),status=422)

    def test_header_and_body_boundary_crossing_rejected(self):
        self.register(merged_source(),header_end=5,status=422)
        self.register(merged_source(),data_end=7,status=422)
        self.register(merged_source(),data_start=6,status=422)

    def test_no_phantom_records_from_merged_only_blank_rows(self):
        def extend(ws):ws.unmerge_cells('D7:D8');ws.merge_cells('D7:D10')
        j=self.upload(self.setup_job(),[],raw=merged_source(names=['가상하나','가상둘'],variant=extend),dep='')
        self.assertEqual(len(j['rows']),2);self.assertEqual(j['rows'][1]['merge_provenance'][0]['range'],'D7:D10')

    def test_formula_in_merged_anchor_rejected(self):
        j=self.setup_job();self.upload(j,[],raw=merged_source(names=['가상하나','가상둘'],variant=lambda ws:setattr(ws['D7'],'value','=1+1')),status=422)

    def test_manual_correction_keeps_original_anchor_record(self):
        j=self.upload(self.setup_job(),[],raw=merged_source(names=['가상하나','가상둘']),dep='')
        original=j['rows'][1]['merge_provenance'];values=j['rows'][1]['values'][:];values[6]='원본 대조 후 수정한 비고'
        j=self.decide(j,j['rows'][1],values=values)
        self.assertEqual(j['rows'][1]['merge_provenance'],original);self.assertEqual(j['rows'][1]['original_values'][2],'')

    def test_ambiguous_parent_only_header_rejected(self):
        self.register(merged_source(),header_end=4,status=422)

    def test_saved_v21_single_header_contract_still_parses(self):
        snapshot=dict(sheet_name='기관회신',start_column=2,headers=test_adhoc_templates.HEADERS,header_row=4,data_start=5,data_end=12)
        values=['001','가상기존','총무팀','참석','AI 교육','2026-11-06','기존 비고']
        parsed=adhoc_templates.parse(test_adhoc_templates.source([values]),snapshot)
        self.assertEqual(parsed,[(5,values)])

    def test_one_logical_field_horizontal_merge_retained(self):
        t=self.register(wide_source());self.assertEqual(len(t['config']['headers']),7)
        self.assertEqual(t['config']['column_spans'][-1],{'start':8,'end':9})
        j=self.create(template_id=t['id'],course='공공 AI 실무교육')
        j=self.upload(j,[],raw=wide_source(names=['가상하나','가상둘']),dep='')
        self.assertEqual(j['rows'][1]['values'][6],'병합 부서 회신 2')
        self.assertTrue(any(p['range']=='H8:I8' and p['direction']=='horizontal' for p in j['rows'][1]['merge_provenance']))

    def test_partial_or_unmerged_wide_field_rejected(self):
        t=self.register(wide_source());j=self.create(template_id=t['id'],course='공공 AI 실무교육')
        self.upload(j,[],raw=wide_source(names=['가상하나','가상둘'],variant=lambda ws:ws.unmerge_cells('H8:I8')),status=422)

    def test_single_logical_field_rectangular_merge_needs_vertical_permission(self):
        def block(ws):ws.unmerge_cells('H7:I7');ws.unmerge_cells('H8:I8');ws.merge_cells('H7:I8')
        self.register(wide_source(variant=block),status=422)
        t=self.register(wide_source(variant=block),fill_merged_columns=[2,6])
        j=self.create(template_id=t['id'],course='공공 AI 실무교육')
        j=self.upload(j,[],raw=wide_source(names=['가상하나','가상둘'],variant=block),dep='')
        self.assertEqual(j['rows'][1]['values'][6],j['rows'][0]['values'][6])
        self.assertTrue(any(p['direction']=='block' and p['inherited'] for p in j['rows'][1]['merge_provenance']))

    def test_saved_v22_schema_two_contract_still_parses(self):
        raw=merged_source(names=['가상하나','가상둘']);t=self.register(merged_source())
        snapshot=t['config'].copy();snapshot.pop('column_spans');snapshot['header_schema_version']=2
        result=adhoc_templates.parse(raw,snapshot)
        self.assertEqual(len(result),2);self.assertEqual(result[1][1][2],'총무팀')

    def test_overlapping_merge_ranges_rejected(self):
        self.register(merged_source(variant=lambda ws:ws.merge_cells('C4:E4')),status=422)
        j=self.setup_job()
        self.upload(j,[],raw=merged_source(names=['가상하나','가상둘'],variant=lambda ws:ws.merge_cells('D8:D9')),status=422)


if __name__=='__main__':unittest.main()
