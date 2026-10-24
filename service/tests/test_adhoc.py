import io
import tempfile
import unittest
import zipfile
from pathlib import Path
from datetime import datetime
from unittest.mock import patch
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook,load_workbook
from app import adhoc


class AdhocTest(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.override=patch.object(adhoc,'DB_PATH',Path(self.tmp.name)/'adhoc.db');self.override.start();adhoc.initialize()
        app=FastAPI();app.include_router(adhoc.router);self.client=TestClient(app)
        self.settings=dict(situation='각 부서의 교육 참석 희망자 취합',title='[합성] 교육 취합',course='AI 교육',event_date='2026-11-06',deadline='2026-10-30',departments=['총무팀','행정팀','안전팀'])

    def tearDown(self):self.client.close();self.override.stop();self.tmp.cleanup()
    def req(self,method,path,body=None,status=200):
        r=self.client.request(method,'/api/adhoc'+path,json=body);self.assertEqual(r.status_code,status,r.text);return r.json()
    def create(self,**changes):return self.req('POST','/jobs',self.settings|changes,201)
    def row(self,dep='총무팀',name='가상직원',wish='희망'):return [dep,name,'주무관','AI 교육','2026-11-06',wish]
    def upload(self,j,rows,dep='총무팀',status=201,raw=None,name='reply.xlsx'):
        raw=raw if raw is not None else adhoc.workbook(j['settings'],rows)
        r=self.client.post(f"/api/adhoc/jobs/{j['id']}/files",data=dict(version=j['version'],department=dep),files={'file':(name,raw,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
        self.assertEqual(r.status_code,status,r.text);return r.json()
    def decide(self,j,row,decision='include',values=None):
        return self.req('PATCH',f"/jobs/{j['id']}/rows/{row['id']}",dict(version=j['version'],values=values or row['values'],decision=decision,reason='원본과 대조한 검토 사유'))
    def empty(self,j,dep):return self.req('POST',f"/jobs/{j['id']}/empty-replies",dict(version=j['version'],department=dep,reason='담당자로부터 희망자 없음을 확인'))
    def clean(self,j):return self.req('POST',f"/jobs/{j['id']}/review-clean",dict(version=j['version'],reason='전체 정상 행을 원본과 대조'))
    def ready(self):
        j=self.upload(self.create(),[self.row(),self.row(name='미희망직원',wish='미희망')]);j=self.clean(j)
        return self.empty(self.empty(j,'행정팀'),'안전팀')

    def test_create_request_template_and_persistence(self):
        j=self.create();self.assertIn('2026-10-30',j['request_draft']);self.assertFalse(j['can_finalize'])
        r=self.client.get(f"/api/adhoc/jobs/{j['id']}/template.xlsx");self.assertEqual(r.status_code,200)
        wb=load_workbook(io.BytesIO(r.content));self.assertEqual([c.value for c in wb['참석희망자'][1]],adhoc.HEADERS);self.assertEqual(wb['참석희망자'].freeze_panes,'A2');wb.close()
        adhoc.initialize();self.assertEqual(self.req('GET',f"/jobs/{j['id']}")['settings'],j['settings'])

    def test_working_report_derived_sheets_and_version_guard(self):
        j=self.upload(self.create(title='=1+1'),[self.row(name='가상인원'),self.row(name='보류인원')])
        j=self.decide(j,j['rows'][0])
        before=self.req('GET',f"/jobs/{j['id']}")
        r=self.client.get(f"/api/adhoc/jobs/{j['id']}/working.xlsx?version={j['version']}")
        self.assertEqual(r.status_code,200)
        wb=load_workbook(io.BytesIO(r.content))
        self.assertEqual(wb['제출용 명단'].max_row,2)
        self.assertEqual(wb['제출용 명단']['C2'].value,'가상인원');self.assertEqual(wb['업무정보']['B2'].value,'=1+1');self.assertEqual(wb['업무정보']['B2'].data_type,'s')
        self.assertEqual(wb['부서별 집계']['B2'].value,1);self.assertEqual(wb['미회신 목록'].max_row,3)
        self.assertIn('작성 중',wb['업무정보']['B7'].value);wb.close()
        self.assertEqual(before,self.req('GET',f"/jobs/{j['id']}"))
        self.assertEqual(self.client.get(f"/api/adhoc/jobs/{j['id']}/working.xlsx?version=1").status_code,409)

    def test_same_form_merge_provenance_and_read_only(self):
        j=self.upload(self.create(),[self.row()]);j=self.upload(j,[self.row('행정팀','가상둘')],'행정팀')
        self.assertEqual(j['summary']['total'],2);self.assertEqual(j['summary']['pending'],2)
        self.assertEqual([r['source_row'] for r in j['rows']],[2,2]);self.assertEqual(len({r['reply_id'] for r in j['rows']}),2)
        before=self.req('GET',f"/jobs/{j['id']}");self.assertEqual(before,self.req('GET',f"/jobs/{j['id']}"))

    def test_same_file_retry_no_duplicate_or_version_bump(self):
        j=self.create();raw=adhoc.workbook(j['settings'],[self.row()]);new=self.upload(j,[],raw=raw)
        again=self.upload(j,[],raw=raw);self.assertEqual(new,again)
        self.upload(new,[],dep='행정팀',raw=raw,status=409)

    def test_different_format_and_formula_atomic_rejection(self):
        j=self.create();self.upload(j,[],raw=b'bad xlsx',status=422)
        wb=Workbook();wb.active.title='참석희망자';wb.active.append(['성명','부서']);raw=io.BytesIO();wb.save(raw);wb.close();self.upload(j,[],raw=raw.getvalue(),status=422)
        wb=load_workbook(io.BytesIO(adhoc.workbook(j['settings'],[self.row()])));wb['참석희망자']['B2']='=1+1';raw=io.BytesIO();wb.save(raw);wb.close();self.upload(j,[],raw=raw.getvalue(),status=422)
        self.assertEqual(self.req('GET',f"/jobs/{j['id']}")['version'],1)

    def test_missing_and_mismatch_must_fix_before_include(self):
        j=self.upload(self.create(),[self.row(name='')]);r=j['rows'][0];self.assertIn('필수값 누락',r['issues'])
        body=dict(version=j['version'],values=r['values'],decision='include',reason='대조 확인 사유입니다')
        self.req('PATCH',f"/jobs/{j['id']}/rows/{r['id']}",body,422)
        j=self.decide(j,r,values=self.row());self.assertEqual(j['rows'][0]['original_values'][1],'');self.assertEqual(j['rows'][0]['values'][1],'가상직원')

    def test_duplicates_not_deleted_or_auto_included(self):
        j=self.upload(self.create(),[self.row(),self.row(),self.row(name='정상직원')]);self.assertEqual(j['summary']['total'],3)
        self.assertEqual(j['summary']['clean_pending'],1);j=self.clean(j);self.assertEqual(j['summary']['pending'],2)
        self.assertTrue(any('중복 의심' in i for i in j['rows'][0]['issues']))
        j=self.decide(j,j['rows'][1],decision='exclude');self.assertFalse(j['rows'][0]['issues'])
        j=self.clean(j);self.assertEqual(j['summary']['attendees'],2)

    def test_empty_reply_distinct_conflict_and_cancel(self):
        j=self.empty(self.create(),'총무팀');self.assertEqual(j['departments'][0]['status'],'희망자 없음')
        j=self.upload(j,[self.row()]);self.assertEqual(j['departments'][0]['status'],'회신 충돌');self.assertFalse(j['can_finalize'])
        j=self.req('POST',f"/jobs/{j['id']}/empty-replies/총무팀/cancel",dict(version=j['version']))
        self.assertEqual(j['departments'][0]['status'],'검토 필요');self.assertEqual(j['summary']['missing'],2)

    def test_finalization_requires_all_replies_and_review(self):
        j=self.upload(self.create(),[self.row()]);self.req('POST',f"/jobs/{j['id']}/finalize",dict(version=j['version']),422)
        j=self.clean(j);self.req('POST',f"/jobs/{j['id']}/finalize",dict(version=j['version']),422)
        j=self.empty(self.empty(j,'행정팀'),'안전팀');self.assertTrue(j['can_finalize'])

    def test_final_snapshot_include_wishes_only_audit_and_immutable(self):
        j=self.ready();final=self.req('POST',f"/jobs/{j['id']}/finalize",dict(version=j['version']),201);e=final['exports'][0]
        url=f"/api/adhoc/jobs/{j['id']}/exports/{e['id']}";raw=self.client.get(url).content
        wb=load_workbook(io.BytesIO(raw));self.assertEqual(wb['참석희망자'].max_row,2)
        self.assertEqual(wb['검토이력'].max_row,3);self.assertEqual(wb['출처'].cell(2,3).value,j['files'][0]['sha256']);wb.close()
        updated=self.decide(j,j['rows'][0],decision='exclude');self.assertFalse(updated['exports'][0]['is_current']);self.assertEqual(raw,self.client.get(url).content)

    def test_stale_version_and_cross_job_row_rejected(self):
        j=self.create();new=self.upload(j,[self.row()]);self.req('POST',f"/jobs/{j['id']}/review-clean",dict(version=j['version'],reason='전체 행 검토했습니다'),409)
        other=self.create();r=new['rows'][0];self.req('PATCH',f"/jobs/{other['id']}/rows/{r['id']}",dict(version=1,values=r['values'],decision='exclude',reason='업무가 다른 행 확인'),404)
        self.req('GET','/jobs/missing',status=404)

    def test_formula_injection_exported_as_literal(self):
        j=self.upload(self.create(),[self.row(name='=HYPERLINK("http://example.invalid")')]);j=self.decide(j,j['rows'][0],decision='exclude')
        raw=adhoc.workbook(j['settings'],[],j);wb=load_workbook(io.BytesIO(raw),data_only=False)
        self.assertEqual(wb['검토이력']['B2'].data_type,'s');wb.close()

    def test_excel_date_and_oversize_empty_file(self):
        j=self.create();row=self.row();row[4]=datetime(2026,11,6)
        wb=Workbook();wb.active.title='참석희망자';wb.active.append(adhoc.HEADERS);wb.active.append(row);s=io.BytesIO();wb.save(s);wb.close()
        j=self.upload(j,[],raw=s.getvalue());self.assertEqual(j['rows'][0]['values'][4],'2026-11-06')
        self.upload(j,[],status=422);self.upload(j,[],raw=b'x'*(adhoc.MAX_BYTES+1),status=422)

    def test_demo_zip_three_departments_and_bad_format(self):
        j=self.create();response=self.client.get(f"/api/adhoc/jobs/{j['id']}/demo.zip")
        with zipfile.ZipFile(io.BytesIO(response.content)) as z:
            names=[n for n in z.namelist() if n.endswith('.xlsx')];self.assertEqual(len(names),4)
            for i,name in enumerate(names[:3]):j=self.upload(j,[],j['settings']['departments'][i],raw=z.read(name),name=name)
            self.upload(j,[],raw=z.read(names[-1]),status=422)
        self.assertEqual(j['summary']['total'],6);self.assertEqual(j['summary']['clean_pending'],3)

    def test_invalid_job_settings(self):
        for changes in [dict(departments=['총무팀','총무팀']),dict(deadline='2026-11-07'),dict(departments=[' '])]:self.req('POST','/jobs',self.settings|changes,422)

    def test_finalization_retry_reuses_current_snapshot(self):
        j=self.ready();path=f"/jobs/{j['id']}/finalize"
        first=self.req('POST',path,dict(version=j['version']),201)
        second=self.req('POST',path,dict(version=j['version']),201)
        self.assertEqual(first,second)
        self.assertEqual(len(second['exports']),1)

    def test_automatic_department_detection_and_ambiguous_rejection(self):
        j=self.upload(self.create(),[self.row('행정팀')],dep='')
        self.assertEqual(j['files'][0]['department'],'행정팀')
        self.assertEqual(j['summary']['clean_pending'],1)
        self.upload(j,[self.row('행정팀'),self.row('안전팀')],dep='',status=422)

    def test_file_department_correction_preserves_original_resets_review(self):
        j=self.upload(self.create(),[self.row()]);j=self.clean(j);f=j['files'][0]
        path=f"/jobs/{j['id']}/files/{f['id']}"
        original=self.client.get('/api/adhoc'+path).content
        j=self.req('PATCH',path,dict(version=j['version'],department='행정팀',reason='회신 부서를 다시 확인한 정정'))
        self.assertEqual(j['rows'][0]['decision'],'pending');self.assertTrue(j['rows'][0]['issues'])
        self.assertEqual(original,self.client.get('/api/adhoc'+path).content)


if __name__=='__main__':unittest.main()
