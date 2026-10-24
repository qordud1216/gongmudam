"""핵심 흐름 시험. 실행: .venv\\Scripts\\python -m unittest discover -s tests -v
임시 폴더에 DB를 만들어 쓰므로 실제 data/ 폴더는 건드리지 않는다."""
from __future__ import annotations

import hashlib
import json
import re
import shutil
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from app import assistant, config, deadlines, formfill, llm, seed, store, travel_api, adhoc, folders  # noqa: E402

TODAY = date(2026, 10, 19)  # 월요일


def _use_temp_store(tmp: Path) -> None:
    store._conn = None
    store.DATA = tmp / "data"
    store.LIBRARY = store.DATA / "library"
    store.DB_PATH = store.DATA / "seomu.db"
    travel_api.DB_PATH = store.DATA / "travel.db"
    adhoc.DB_PATH = store.DATA / "adhoc.db"
    folders.DB_PATH = store.DATA / "folders.db"
    formfill.OUTPUT = tmp / "output"
    config.CONFIG_PATH = tmp / "config.json"


def _sample_hashes() -> dict[str, str]:
    return {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in (ROOT / "sample_data").rglob("*") if p.is_file()}


class FakeProvider(llm.BaseProvider):
    """모델 응답을 흉내 낸다. 실제 모델 없이 LLM 경로(인용 검증·JSON 해석)를 시험한다."""
    reply = ""

    def chat(self, messages, **kw):
        return FakeProvider.reply


class Base(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = Path(tempfile.mkdtemp())
        _use_temp_store(cls.tmp)
        cls.before = _sample_hashes()
        seed.load_seed(reset=True)

    @classmethod
    def tearDownClass(cls):
        if store._conn:
            store._conn.close()
            store._conn = None
        shutil.rmtree(cls.tmp, ignore_errors=True)


class DeadlineTest(unittest.TestCase):
    def test_business_days_skip_weekend_and_holiday(self):
        # 2026-10-08(목) + 1업무일 → 10-09 한글날 건너뛰고 10-12(월)
        self.assertEqual(deadlines.add_business_days(date(2026, 10, 8), 1), date(2026, 10, 12))
        # 금요일 + 3업무일 → 다음 주 수요일
        self.assertEqual(deadlines.add_business_days(date(2026, 10, 16), 3), date(2026, 10, 21))

    def test_calendar_rule_moves_before_holiday(self):
        # 행사일 2026-10-24(토) - 7일 = 10-17(토) → 휴일이면 앞당겨 10-16(금)
        due = deadlines.apply_rule({"base": "행사일", "offset": -7, "unit": "일"}, {"행사일": date(2026, 10, 24)})
        self.assertEqual(due, date(2026, 10, 16))

    def test_find_dates(self):
        found = [f["date"] for f in deadlines.find_dates("어제 출장 다녀왔고 11월 20일에 행사", TODAY)]
        self.assertEqual(found, [date(2026, 10, 18), date(2026, 11, 20)])
        nxt = deadlines.find_dates("다음 주 월요일 연가", TODAY)[0]["date"]
        self.assertEqual(nxt, date(2026, 10, 26))
        # '시작일' 속 '작일', '설명일' 속 '명일'은 상대 날짜가 아니다
        self.assertEqual(deadlines.find_dates("연가 시작일과 설명일", TODAY), [])
        self.assertEqual(deadlines.find_dates("명일 제출", TODAY)[0]["date"], date(2026, 10, 20))


class AskOfflineTest(Base):
    def setUp(self):
        cfg = config.load()
        cfg["llm"]["active"] = "offline"
        config.save(cfg)

    def test_travel_settlement(self):
        r = assistant.ask("어제 서울 출장 다녀왔어요. 정산은 언제까지 하나요?", today=TODAY)
        self.assertEqual(r["mode"], "offline")
        p = r["plan"]
        self.assertEqual(p["card"]["id"], "card-travel")
        self.assertEqual(p["base_values"], {"출장 종료일": "2026-10-18"})
        step3 = p["steps"][2]
        self.assertEqual(step3["due"], "2026-10-21")  # 일요일 종료 → 다음 업무일부터 3업무일
        self.assertEqual(step3["forms"][0]["id"], "form-travel")
        labels = [b["label"] for b in step3["basis"]]
        self.assertIn("제3조(출장 여비 정산 기한)", labels)
        self.assertTrue(all(b.get("text") for b in step3["basis"]))
        self.assertIn("지금 할 일", r["summary"])
        self.assertEqual(r["answer"], "")
        self.assertIsNone(r["ai_note"])

    def test_range_sets_both_bases(self):
        r = assistant.ask("10월 13일부터 15일까지 출장 다녀왔어요", today=TODAY)
        self.assertEqual(r["plan"]["base_values"], {"출장 시작일": "2026-10-13", "출장 종료일": "2026-10-15"})

    def test_other_cards(self):
        cases = {"11월 20일에 어린이 안전체험 행사를 열어요": "card-event",
                 "소화기 어제 납품받았어요": "card-purchase",
                 "다음 주 월요일에 연가 쓰려고요": "card-leave"}
        for q, cid in cases.items():
            with self.subTest(q=q):
                self.assertEqual(assistant.ask(q, today=TODAY)["plan"]["card"]["id"], cid)

    def test_draft_card_warns(self):
        p = assistant.ask("다음 주 월요일에 연가 쓰려고요", today=TODAY)["plan"]
        self.assertTrue(any("초안" in w for w in p["warnings"]))

    def test_unknown_question(self):
        r = assistant.ask("점심 메뉴 추천해 줘", today=TODAY)
        self.assertIsNone(r["plan"])
        self.assertIn("찾지 못했습니다", r["summary"])

    def test_december_travel_cap(self):
        """F1: 12월에 끝난 출장은 기관 기준 제3조②에 따라 12월 26일보다 늦지 않게 안내한다."""
        r = assistant.ask("12월 22일 서울 출장 다녀왔어요. 정산 언제까지 하나요?", today=date(2026, 12, 23))
        p = r["plan"]
        self.assertEqual(p["base_values"], {"출장 종료일": "2026-12-22"})
        step3 = p["steps"][2]
        self.assertLessEqual(step3["due"], "2026-12-26")  # 고정 +3업무일이면 12-28
        self.assertEqual(step3["due"], "2026-12-24")  # 12-26(토)·12-25(성탄절) → 앞 업무일
        self.assertTrue(any("제3조②" in w and "12월 26일" in w for w in p["warnings"]))
        self.assertIn("제3조(출장 여비 정산 기한)", [b["label"] for b in step3["basis"]])
        # 12월이 아니면 그대로 3업무일
        p = assistant.ask("11월 20일 출장 다녀왔어요", today=date(2026, 11, 20))["plan"]
        self.assertEqual(p["steps"][2]["due"], "2026-11-25")
        self.assertFalse(any("제3조②" in w for w in p["warnings"]))
        # 26일 이후에 끝난 출장은 기한이 이미 지났다고 알린다
        p = assistant.ask("12월 29일 출장 끝났어요 정산", today=date(2026, 12, 29))["plan"]
        self.assertEqual(p["steps"][2]["dday"]["state"], "overdue")
        self.assertTrue(any("이미 지났습니다" in w for w in p["warnings"]))

    def test_shift_worker_leave(self):
        """F2: 교대근무자는 연가 안내 제3조에 따라 7일 전까지. 조건이 없으면 해당자 기한을 함께 알린다."""
        p = assistant.ask("교대근무자인데 다음 주 월요일에 연가 쓰려고요", today=TODAY)["plan"]
        self.assertEqual(p["base_values"], {"연가 시작일": "2026-10-26"})
        self.assertEqual(p["conditions"], ["shift"])
        self.assertEqual(p["steps"][0]["due"], "2026-10-19")
        self.assertEqual(p["steps"][0]["variant"], "교대근무자")
        self.assertTrue(any("교대근무자' 조건을 적용" in w for w in p["warnings"]))
        p = assistant.ask("다음 주 월요일에 연가 쓰려고요", today=TODAY)["plan"]
        self.assertEqual(p["steps"][0]["due"], "2026-10-23")
        self.assertTrue(any("교대근무자: 2026. 10. 19.(월)까지" in w for w in p["warnings"]))
        # 홈 등록·재계산(/api/plan)에서도 같은 조건을 넘기면 같은 기한이 나온다
        card = store.get_card("card-leave")
        p = assistant.build_plan(card, {"연가 시작일": date(2026, 10, 26)}, TODAY, ["shift"])
        self.assertEqual(p["steps"][0]["due"], "2026-10-19")
        # 질병·사고로 미리 못 낸 경우는 제2조 절차로 안내한다
        p = assistant.ask("어제 아파서 연가를 못 냈어요", today=TODAY)["plan"]
        self.assertEqual(p["conditions"], ["urgent"])
        self.assertIn("복귀한 날", p["steps"][0]["detail"])

    def test_compound_conditions(self):
        """R1: 교대근무와 질병·사고가 함께 해당되면 긴급 예외(제2조)를 우선하고, 교대 기한(10-12)을 단정하지 않는다."""
        q = "교대근무자인데 어제 사고로 연가를 미리 신청하지 못했어요. 오늘 복귀했어요"
        p = assistant.ask(q, today=TODAY)["plan"]
        self.assertEqual(p["conditions"], ["shift", "urgent"])
        st = p["steps"][0]
        self.assertEqual(st["variant"], "질병·사고로 미리 신청하지 못한 경우")
        self.assertIsNone(st["due"])
        self.assertIn("복귀한 날", st["detail"])
        self.assertTrue(any("'교대근무자' 조건도 해당되지만" in w and "적용하지 않았습니다" in w for w in p["warnings"]))
        # 홈 등록·재계산으로 두 조건을 함께 넘겨도 결과가 같다(순서와 무관)
        card = store.get_card("card-leave")
        for cond in (["shift", "urgent"], ["urgent", "shift"]):
            p2 = assistant.build_plan(card, {"연가 시작일": date(2026, 10, 18)}, TODAY, cond)
            self.assertIsNone(p2["steps"][0]["due"])
            self.assertEqual(p2["steps"][0]["variant"], st["variant"])
        # 우선순위를 정하지 않은 조건끼리 겹치면 어느 기한도 단정하지 않고 충돌을 알린다
        tie = {"id": "t", "title": "시험", "category": "", "layer": "기관", "summary": "", "status": "확정", "notes": "",
               "base_dates": ["시작일"], "steps": [{"title": "신청", "deadline": {"base": "시작일", "offset": -1, "unit": "업무일"},
               "variants": [{"id": "a", "label": "가", "deadline": {"base": "시작일", "offset": -7, "unit": "일"}},
                            {"id": "b", "label": "나", "deadline": {"base": "시작일", "offset": -3, "unit": "일"}}]}]}
        p3 = assistant.build_plan(tie, {"시작일": date(2026, 10, 30)}, TODAY, ["a", "b"])
        self.assertIsNone(p3["steps"][0]["due"])
        self.assertTrue(any("기한을 하나로 정할 수 없습니다" in w for w in p3["warnings"]))
        # 모의 절차 카드 전체: 한 단계에 조건이 여러 개면 모두 감지돼도 기한이 하나로 정해지거나 충돌로 표시된다
        for c in store.list_cards():
            for i, s in enumerate(c["steps"]):
                ids = [v["id"] for v in s.get("variants") or []]
                if len(ids) < 2:
                    continue
                pc = assistant.build_plan(c, {b: date(2026, 10, 26) for b in c["base_dates"]}, TODAY, ids)
                stp = pc["steps"][i]
                self.assertTrue(stp["variant"] or any("하나로 정할 수 없습니다" in w for w in pc["warnings"]), c["id"])

    def test_layer_filter(self):
        r = assistant.ask("출장 정산", today=TODAY, layers=["공통"])
        self.assertIsNone(r["plan"])  # 출장 카드는 기관 층
        self.assertTrue(all(c["layer"] == "공통" for c in r["related"]))


class LLMPathTest(Base):
    def setUp(self):
        self._orig = dict(llm.PROVIDERS)
        llm.PROVIDERS["openai_compatible"] = FakeProvider
        cfg = config.load()
        self._active = cfg["llm"]["active"]
        cfg["llm"]["active"] = "local_vllm"
        config.save(cfg)

    def tearDown(self):
        llm.PROVIDERS.clear()
        llm.PROVIDERS.update(self._orig)
        cfg = config.load()
        cfg["llm"]["active"] = self._active
        config.save(cfg)

    def test_fake_citation_is_dropped(self):
        q = "지난주 금요일까지 출장이었어요 정산 어떻게 해요"
        real = assistant.search_chunks(q)[0][0]["id"]  # 모델에게 실제로 준 후보 조문
        FakeProvider.reply = ('```json\n{"card_id": "card-travel", "answer": "정산서를 3업무일 안에 내세요.", '
                              f'"cited": ["{real}", "없는-조문-99"], "dates": {{"출장 종료일": "2026-10-16"}}}}\n```')
        r = assistant.ask(q, today=TODAY)
        self.assertEqual(r["mode"], "llm")
        self.assertIn("존재하지 않는 근거 1건", r["notice"])
        self.assertEqual([c["id"] for c in r["cited"]], [real])
        self.assertEqual(r["plan"]["card"]["id"], "card-travel")

    def _ask_with(self, answer: str, q: str = "어제 서울 출장 다녀왔어요. 정산은 언제까지 하나요?",
                  today: date = TODAY, card: str = "card-travel") -> dict:
        FakeProvider.reply = json.dumps({"card_id": card, "answer": answer, "cited": []}, ensure_ascii=False)
        return assistant.ask(q, today=today)

    # 기준일 10-19(월), 출장 종료 10-18 → 정산서 기한 10-21(수). 모델 설명에 넣어 보는 기한·날짜 문장들.
    SCHEDULE_SENTENCES = [
        # Codex 4차 재현
        "정산서 제출일은 시월 스무날입니다.",
        "정산서 제출일은 Tue입니다.",
        "정산서 제출은 Tuesday입니다.",
        "정산서는 D−2입니다.",
        "정산서는 차 주에 제출하세요.",
        # Codex 3차 재현
        "정산서 제출 기한은 오는 20일에 도래합니다.",
        "정산서 제출 마감은 하루 뒤입니다.",
        "정산서 제출 마감일은 20일로 정해져 있습니다.",
        "정산서 제출 마감은 10월 21일(목)까지입니다.",
        "오늘까지 증빙서류를 모아 국내출장 여비 정산서를 작성하여 제출하세요.",
        "사용하려는 날의 전 업무일까지 신청 규정을 준수하세요.",
        "정산서는 오늘까지, 증빙은 10월 21일까지 내세요.",
        "명일이 아니라 모레인 10월 21일까지 내세요.",
        # 이전 라운드 재현과 변형
        "정산서 제출 기한은 내일인 10월 21일까지입니다.",
        "내일 (10월 21일) 까지 여비 정산서를 제출하세요.",
        "정산서는 내일 중 제출하세요.",
        "정산서를 내일자로 제출하세요.",
        "10월 21일(내일)에 내세요.",
        "정산서는 모레 내세요.",
        "정산서는 사흘 안에 내세요.",
        "정산서는 익일 제출하세요.",
        "정산서는 차주에 내세요.",
        "정산서는 이번 주 중으로 내세요.",
        "정산서는 다음 달 초에 내세요.",
        "정산서는 삼일 안에 내세요.",
        "정산서 제출은 D-2입니다.",
        "정산서 제출 D-day를 확인하세요.",
        "정산서는 수요일에 내세요.",
        "정산서는 (수)에 내세요.",
        "정산서는 2026. 10. 21.에 내세요.",
        "정산서는 10/21에 내세요.",
        "정산서는 3업무일 안에 내세요.",
        "정산서는 늦어도 수요일 오전에 내세요.",
        "정산서는 출장에서 돌아온 즉시 내세요.",
        "정산서는 지체 없이 제출하세요.",
        "정산서는 작일 기준으로 계산합니다.",
        "정산서는 금일 중 제출하세요.",
        "정산서는 주말 전에 내세요.",
        "정산서는 일주일 안에 내세요.",
        "정산서는 열흘 안에 내세요.",
        "정산서는 내 일 낼 수 있습니다.",
        "정산서는 내​일 내세요.",
        "정산서는 ２１일에 내세요.",
        "정산서는 출장에서 돌아온 날 내세요.",
        "Submit it by tomorrow.",
        "정산서는 21日에 내세요.",
        "정산서는 석 달 안에 내세요.",
        # 4차 보강: 고유어 날짜·한글 월명·기호·영문·띄어쓰기 변형
        "정산서는 보름 안에 내세요.",
        "정산서는 그믐에 내세요.",
        "정산서는 초사흘에 내세요.",
        "정산서는 스무하루에 내세요.",
        "정산서는 열하루에 내세요.",
        "정산서는 동짓달에 내세요.",
        "정산서는 섣달 안에 내세요.",
        "정산서는 유월에 내세요.",
        "정산서는 이달 하순에 내세요.",
        "정산서는 마 감 전에 내세요.",
        "정산서는 오 늘 내세요.",
        "정산서는 다 음 주 에 내세요.",
        "정산서는 D–2입니다.",
        "정산서는 D－2입니다.",
        "정산서는 Wed에 내세요.",
        "정산서는 ASAP 제출하세요.",
        "정산서는 by EOD 제출하세요.",
        "정산서는 木曜에 내세요.",
        "정산서는 #2 순서로 내세요.",
        "정산서는 ① 단계에서 내세요.",
        "정산서 제출일은 이십일일입니다.",
        "정산서 제출은 익영업일입니다.",
        "정산서는 두 번째 업무일에 제출하세요.",
        "정산서를 미리 준비하세요.",
        "정산서를 제때 제출하세요.",
        "증빙을 준비하세요.정산서는 D−2입니다.",
    ]
    # 7차: '에·에서' 조사, 부사('먼저·함께·실제'), '사전·날짜', 연결 어미('-고')가 없는 절차 문장만 남는다.
    PLAIN_SENTENCES = [
        "국내출장 여비 정산서의 영수증을 확인하세요.",
        "기차 승차권과 숙박 영수증을 증빙으로 첨부하세요.",
        "출장 결재 문서 번호를 적으세요.",
        "기관 지침을 따르세요.",
        "운임은 지출한 금액을 적습니다.",
        "증빙이 없는 비용은 빠집니다.",
        "등록된 규정으로 확인되지 않는 비용은 담당 부서로 문의하세요.",
        "부서장 승인 문서를 첨부하세요.",
        "서식은 서식 화면으로 내려받을 수 있습니다.",
        "전자결재로 승인 여부를 확인하세요.",
        "정산서의 서명 날인을 확인하세요.",
    ]
    # 6차까지 남던 문장. 7차 허용 목록(조사 '에', 부사, '사전·날짜', 연결 어미 제외)에서는 문장째 빠진다(정상 설명 감소 감수).
    R7_DROPPED_PLAIN = [
        "국내출장 여비 정산서에 영수증을 붙여 결재를 올리세요.",
        "출장 결재 문서 번호를 정산서에 적으세요.",
        "기관 지침이 공통 규정보다 먼저 적용됩니다.",
        "운임은 실제 지출한 금액만 적습니다.",
        "증빙이 없는 비용은 정산 대상에서 빠집니다.",
        "부서장 승인 문서를 함께 첨부하세요.",
        "전자결재로 올리고 사전 승인 여부를 확인하세요.",
        "정산서에 서명 날인하고 작성 날짜를 적으세요.",
    ]

    def test_answer_never_uses_model_text(self):
        """R2 구조: 본문(answer)은 규칙 기반 문장이라 모델이 무엇을 쓰든 같은 카드면 한 글자도 바뀌지 않는다."""
        self.assertGreaterEqual(len(self.SCHEDULE_SENTENCES), 60)
        base = self._ask_with("")
        expected = base["answer"]
        dues = [st["due"] for st in base["plan"]["steps"]]
        self.assertEqual(dues[2], "2026-10-21")
        self.assertEqual(expected, assistant._rule_answer(base["plan"]))
        for s in ["정산서", "아무 말", "{}", "\n\n", "x" * 500, *self.SCHEDULE_SENTENCES, *self.PLAIN_SENTENCES,
                  " ".join(self.SCHEDULE_SENTENCES), " ".join(self.PLAIN_SENTENCES)]:
            r = self._ask_with(s)
            self.assertEqual(r["mode"], "llm", s)
            self.assertEqual(r["answer"], expected, s)
            self.assertEqual([st["due"] for st in r["plan"]["steps"]], dues, s)
            self.assertEqual(r["summary"], base["summary"], s)

    def test_rule_answer_content(self):
        """본문은 단계 이름·필요 서식·근거 조문·적용 조건으로 만든다."""
        r = self._ask_with("")
        a = r["answer"]
        for st in r["plan"]["steps"]:
            self.assertIn(f"{st['no']}단계 {st['title']}", a)
        self.assertIn("국내출장 여비 정산서", a)
        self.assertIn("제3조(출장 여비 정산 기한)", a)
        self.assertNotIn("적용한 조건", a)
        r = self._ask_with("", q="교대근무자인데 어제 사고로 연가를 미리 신청하지 못했어요. 오늘 복귀했어요", card="card-leave")
        self.assertIn(f"적용한 조건은 {r['plan']['steps'][0]['variant']}입니다.", r["answer"])
        self.assertEqual(assistant._rule_answer(None), "관련 규정 조문을 확인하세요.")

    def test_ai_note_drops_schedule_sentences(self):
        """AI 보충 설명에서는 기한·날짜를 말할 수 있는 문장을 날짜 값과 상관없이 문장째 뺀다."""
        keep = "결재 문서 번호를 적으세요."
        for s in self.SCHEDULE_SENTENCES:
            r = self._ask_with(f"{s} {keep}")
            self.assertEqual(r["ai_note"], keep, s)
            self.assertEqual(r["ai_note_dropped"], 1, s)
            self.assertIsNone(r["notice"], s)
            r = self._ask_with(s)
            self.assertIsNone(r["ai_note"], s)  # 모두 빠지면 영역을 숨긴다
            self.assertEqual(r["ai_note_dropped"], 0, s)
        # 줄바꿈으로 나눈 목록도 줄마다 판단한다(번호 '1.'처럼 숫자가 든 줄도 뺀다)
        r = self._ask_with("증빙을 모으세요\n10월 21일까지 정산서를 내세요\n" + keep)
        self.assertEqual(r["ai_note"], "증빙을 모으세요 " + keep)
        self.assertEqual(r["ai_note_dropped"], 1)

    def test_ai_note_keeps_plain_procedure(self):
        """날짜 낱말·숫자·영문이 없는 절차·서식·근거 설명은 보충 설명에 남는다."""
        for s in self.PLAIN_SENTENCES:
            r = self._ask_with(s)
            self.assertEqual(r["ai_note"], s)
            self.assertEqual(r["ai_note_dropped"], 0)
        r = self._ask_with(" ".join(self.PLAIN_SENTENCES[:3]))
        self.assertEqual(r["ai_note"], " ".join(self.PLAIN_SENTENCES[:3]))
        # 숫자·영문이 든 문장은 날짜가 아니어도 뺀다(오탐 감수)
        r = self._ask_with("KTX 승차권을 첨부하세요. 기관 지침 제3조를 따르세요. " + self.PLAIN_SENTENCES[0])
        self.assertEqual(r["ai_note"], self.PLAIN_SENTENCES[0])
        self.assertEqual(r["ai_note_dropped"], 2)
        for s in self.R7_DROPPED_PLAIN:
            self.assertIsNone(self._ask_with(s)["ai_note"], s)

    def test_rule_warnings_unchanged_with_llm(self):
        """모델 문장과 상관없이 계산 카드와 규칙 경고(R1 우선순위·F2 교대)는 그대로다."""
        r = self._ask_with("교대근무자는 7일 전인 오늘까지 신청하세요. 부서장 승인을 받으세요.",
                           q="교대근무자인데 다음 주 월요일에 연가 쓰려고요", card="card-leave")
        self.assertEqual(r["ai_note"], "부서장 승인을 받으세요.")
        self.assertEqual(r["plan"]["steps"][0]["due"], "2026-10-19")
        r = self._ask_with("복귀하신 오늘 오전 중 부서장에게 알리세요. 복무관리시스템으로 신청하세요.",
                           q="교대근무자인데 어제 사고로 연가를 미리 신청하지 못했어요. 오늘 복귀했어요", card="card-leave")
        self.assertEqual(r["ai_note"], "복무관리시스템으로 신청하세요.")
        self.assertIsNone(r["plan"]["steps"][0]["due"])
        self.assertTrue(any("교대근무자" in w for w in r["plan"]["warnings"]))

    def test_broken_reply_falls_back(self):
        FakeProvider.reply = "죄송합니다. 잘 모르겠습니다."
        r = assistant.ask("어제 출장 다녀왔어요", today=TODAY)
        self.assertEqual(r["mode"], "fallback")
        self.assertEqual(r["answer"], "")
        self.assertIsNone(r["ai_note"])
        self.assertEqual(r["plan"]["card"]["id"], "card-travel")

    def test_formfill_rejects_ungrounded_values(self):
        FakeProvider.reply = '{"출장지": "서울", "출장 목적": "정책 회의", "운임": "59800", "숙박비": "150000"}'
        f = store.get_form("form-travel")
        r = formfill.prefill(f, query="어제 서울 정책 회의 출장 다녀왔고 KTX 59,800원 썼어요",
                             base_values={"출장 종료일": "2026-10-18"}, today=TODAY)
        v = r["values"]
        self.assertEqual(v["출장지"], "서울")
        self.assertEqual(v["운임"], "59800")
        self.assertEqual(v["숙박비"], "")  # 문장에 없는 금액은 버린다
        self.assertEqual(v["성명"], "홍길동(가상)")


class RegistryTest(Base):
    def test_crud_and_basis_tracking(self):
        src = self.tmp / "임시규정.txt"
        src.write_text("제1조(목적) 시험용 규정이다.\n제2조(기한) 신청은 5일 전까지 한다.\n", encoding="utf-8")
        before = src.read_bytes()
        d = store.add_document(src, title="시험 규정", kind="규정", layer="개인")
        self.assertEqual(src.read_bytes(), before)  # 원본은 그대로
        self.assertEqual([c["label"] for c in d["chunks"]], ["제1조(목적)", "제2조(기한)"])
        with self.assertRaises(ValueError):
            store.add_document(src, title="중복", kind="규정", layer="개인")

        card = store.save_card({"title": "시험 절차", "category": "시험", "layer": "개인", "status": "확정",
                                "base_dates": ["신청일"],
                                "steps": [{"title": "신청", "deadline": {"base": "신청일", "offset": -5, "unit": "일"},
                                           "basis": [{"doc_id": d["id"], "label": "제2조(기한)"}]}]})
        p = assistant.build_plan(card, {"신청일": date(2026, 11, 2)}, TODAY)
        self.assertEqual(p["steps"][0]["due"], "2026-10-28")
        self.assertFalse(p["warnings"])

        # 개정본으로 교체 → 같은 ID 유지, 확정 이후 수정 경고
        import time
        time.sleep(1.1)
        src2 = self.tmp / "임시규정_개정.txt"
        src2.write_text("제1조(목적) 시험용 규정이다.\n제2조(기한) 신청은 7일 전까지 한다.\n", encoding="utf-8")
        store.add_document(src2, title="시험 규정", kind="규정", layer="개인", doc_id=d["id"])
        p = assistant.build_plan(store.get_card(card["id"]), {"신청일": date(2026, 11, 2)}, TODAY)
        self.assertIn("7일 전", p["steps"][0]["basis"][0]["text"])
        self.assertTrue(any("수정되었습니다" in w for w in p["warnings"]))

        # 문서 삭제 → 근거 없음 표시
        store.delete_document(d["id"])
        p = assistant.build_plan(store.get_card(card["id"]), {}, TODAY)
        self.assertTrue(p["steps"][0]["basis"][0]["missing"])
        self.assertTrue(store.delete_card(card["id"]))

    def test_sample_files_untouched(self):
        seed.load_seed(reset=True)
        self.assertEqual(_sample_hashes(), self.before)

    def test_all_seed_basis_resolve(self):
        for c in store.list_cards():
            for s in c["steps"]:
                for b in s.get("basis") or []:
                    with self.subTest(card=c["id"], label=b["label"]):
                        self.assertIsNotNone(store.resolve_basis(b["doc_id"], b["label"]))

    def test_export_docx(self):
        f = store.get_form("form-travel")
        path = formfill.export_docx(f, {"성명": "홍길동(가상)", "출장지": "서울"}, card_title="국내 출장 여비 정산")
        self.assertTrue(path.exists())
        from docx import Document
        cells = [c.text for row in Document(path).tables[0].rows for c in row.cells]
        self.assertIn("서울", cells)


class ApiTest(Base):
    def test_endpoints(self):
        from fastapi.testclient import TestClient
        from app.main import app
        with TestClient(app) as c:
            self.assertEqual(c.get("/").status_code, 200)
            # This fixture intentionally exercises the synthetic seed-card flow.
            r = c.post("/api/ask", json={"query": "어제 출장 다녀왔어요", "today": "2026-10-19",
                                          "card_id": "card-travel"}).json()
            self.assertEqual(r["plan"]["card"]["id"], "card-travel")
            r = c.post("/api/plan", json={"card_id": "card-event", "base_values": {"행사일": "2026-11-20"}, "today": "2026-10-19"}).json()
            self.assertEqual(r["steps"][0]["due"], "2026-10-30")
            up = c.post("/api/documents", files={"file": ("메모.txt", "제1조(시험) 내용".encode("utf-8"), "text/plain")},
                        data={"title": "API 시험", "kind": "메모", "layer": "개인"})
            self.assertEqual(up.status_code, 200, up.text)
            bad = c.post("/api/documents", files={"file": ("a.exe", b"MZ", "application/octet-stream")}, data={})
            self.assertEqual(bad.status_code, 400)
            did = up.json()["id"]
            self.assertEqual(c.put(f"/api/documents/{did}", json={"title": "이름 바꿈"}).json()["title"], "이름 바꿈")
            self.assertTrue(c.delete(f"/api/documents/{did}").json()["ok"])
            ex = c.post("/api/forms/form-inspect/export", json={"values": {"품명": "소화기"}})
            self.assertEqual(ex.status_code, 200)
            self.assertTrue(ex.content[:2] == b"PK")
            self.assertTrue(c.post("/api/settings/llm/test", json={"profile": "offline"}).json()["ok"])
            s = c.get("/api/settings").json()
            self.assertNotIn("api_key", s["llm"]["profiles"]["openai"])


    def test_home_tracking(self):
        from fastapi.testclient import TestClient
        from app.main import app
        with TestClient(app) as c:
            self.assertEqual(c.get("/api/home?today=2026-10-19").json()["tasks"], [])
            t = c.post("/api/tracked", json={"card_id": "card-event", "base_values": {"행사일": "2026-11-20"},
                                            "label": "어린이 안전체험 행사"}).json()
            h = c.get("/api/home?today=2026-10-19").json()
            self.assertEqual(h["tasks"][0]["next"]["no"], 1)
            self.assertEqual(h["upcoming"][0]["due"], "2026-10-30")
            self.assertEqual([u["due"] for u in h["upcoming"]], sorted(u["due"] for u in h["upcoming"]))
            c.put(f"/api/tracked/{t['id']}", json={"done": [0]})
            h = c.get("/api/home?today=2026-10-19").json()
            self.assertEqual(h["tasks"][0]["next"]["no"], 2)
            self.assertNotIn(1, [u["no"] for u in h["upcoming"]])
            self.assertEqual(c.post("/api/tracked", json={"card_id": "없음"}).status_code, 400)
            self.assertEqual(c.put("/api/tracked/없음", json={"done": []}).status_code, 404)
            c.post("/api/settings/reset-sample")
            self.assertEqual(c.get("/api/home").json()["tasks"], [])



class ServerErrorLeakTest(Base):
    """모델 서버가 오류·잘못된 형식으로 응답해도 그 본문이 화면(notice·summary·warnings 등)에 나가지 않아야 한다.
    실제 OpenAICompatibleProvider 를 쓰고 httpx.post 만 바꿔 끼운다(Codex 5차 중간 보고 재현)."""
    LEAK = "정산서는 오늘 10월 19일까지 제출하세요"
    MARKERS = [LEAK, "10월 19일", "제출하세요"]  # 규칙 기반 summary 는 "오늘까지"(D-day) 를 정상적으로 쓰므로 "오늘"은 표지에서 뺀다

    class _Resp:
        def __init__(self, status: int, text: str):
            self.status_code, self.text = status, text

        def json(self):
            return json.loads(self.text)

    def setUp(self):
        import httpx
        self._post = llm.httpx.post
        cfg = config.load()
        self._active = cfg["llm"]["active"]
        cfg["llm"]["active"] = "local_vllm"
        config.save(cfg)
        self.httpx = httpx

    def tearDown(self):
        llm.httpx.post = self._post
        cfg = config.load()
        cfg["llm"]["active"] = self._active
        config.save(cfg)

    def _serve(self, status: int, text: str):
        llm.httpx.post = lambda *a, **kw: self._Resp(status, text)

    def _cases(self):
        leak = self.LEAK
        chat = lambda content: json.dumps({"choices": [{"message": {"content": content}}]}, ensure_ascii=False)
        return {
            "500 본문": (500, f"Internal Server Error: {leak}"),
            "500 JSON 본문": (500, json.dumps({"error": {"message": leak}}, ensure_ascii=False)),
            "200 + JSON 아님": (200, f"<html>{leak}</html>"),
            "200 + choices 문자열": (200, json.dumps({"choices": leak}, ensure_ascii=False)),
            "200 + 최상위 배열": (200, json.dumps([leak], ensure_ascii=False)),
            "200 + content 객체": (200, json.dumps({"choices": [{"message": {"content": {"x": leak}}}]}, ensure_ascii=False)),
            "200 + content 에 JSON 없음": (200, chat(leak)),
            "200 + content 가 배열 JSON": (200, chat(json.dumps([leak], ensure_ascii=False))),
            "200 + 깨진 JSON": (200, chat('{"card_id": "card-travel", "answer": "' + leak)),
        }

    def _assert_clean(self, r: dict, label: str, keys=("notice", "summary", "answer", "ai_note")):
        for k in keys:
            v = r.get(k) or ""
            for m in self.MARKERS:
                self.assertNotIn(m, v, f"{label}: {k}에 '{m}' 노출 → {v}")
        for w in (r.get("plan") or {}).get("warnings") or []:
            for m in self.MARKERS:
                self.assertNotIn(m, w, f"{label}: warnings에 '{m}' 노출 → {w}")
        self.assertNotIn(self.LEAK, json.dumps(r, ensure_ascii=False), label)

    def test_ask_fallback_hides_server_body(self):
        q = "어제 출장 다녀왔어요 정산 어떻게 해요"
        offline = None
        for label, (status, text) in self._cases().items():
            with self.subTest(label):
                self._serve(status, text)
                with self.assertLogs("app.llm", level="WARNING"):  # 본문은 로그에만 남는다
                    r = assistant.ask(q, today=TODAY)
                self.assertEqual(r["mode"], "fallback", label)
                self.assertIn("모델 서버 응답 오류로 규칙 기반 안내로 전환했습니다", r["notice"])
                self.assertIn("오류 종류:", r["notice"])
                self._assert_clean(r, label)
                self.assertEqual(r["plan"]["card"]["id"], "card-travel")
                if offline is None:
                    offline = r["summary"]
                self.assertEqual(r["summary"], offline, label)  # 규칙 기반 안내는 오류 종류와 무관하게 같다
        self._serve(500, self.LEAK)
        self.assertIn("오류 종류: 모델 서버 HTTP 500 오류", assistant.ask(q, today=TODAY)["notice"])

    def test_connection_error_hides_message(self):
        def boom(*a, **kw):
            raise self.httpx.ConnectError(self.LEAK)
        llm.httpx.post = boom
        r = assistant.ask("어제 출장 다녀왔어요 정산 어떻게 해요", today=TODAY)
        self.assertEqual(r["mode"], "fallback")
        self.assertIn("모델 서버 연결 실패", r["notice"])
        self._assert_clean(r, "연결 실패")

    def test_prefill_hides_server_body(self):
        f = store.get_form("form-travel")
        for label, (status, text) in self._cases().items():
            with self.subTest(label):
                self._serve(status, text)
                r = formfill.prefill(f, query="어제 서울 정책 회의 출장 다녀왔고 KTX 59,800원 썼어요", today=TODAY)
                if label != "200 + content 가 배열 JSON":  # JSON 객체가 아니면 기존대로 조용히 건너뛴다(빈칸 유지)
                    self.assertIn("오류 종류:", r["notice"])
                self._assert_clean(r, label, keys=("notice",))
                self.assertNotIn(self.LEAK, json.dumps(r["values"], ensure_ascii=False))

    def test_settings_llm_test_hides_server_body(self):
        from fastapi.testclient import TestClient
        from app.main import app
        with TestClient(app) as c:
            for label, (status, text) in list(self._cases().items())[:5]:
                with self.subTest(label):
                    self._serve(status, text)
                    r = c.post("/api/settings/llm/test", json={"profile": "local_vllm"}).json()
                    self.assertFalse(r["ok"])
                    self.assertIn("오류 종류:", r["message"])
                    self.assertNotIn(self.LEAK, r["message"])
            # 연결은 됐지만 모델 문장은 화면에 옮기지 않는다
            self._serve(200, json.dumps({"choices": [{"message": {"content": self.LEAK}}]}, ensure_ascii=False))
            r = c.post("/api/settings/llm/test", json={"profile": "local_vllm"}).json()
            self.assertTrue(r["ok"])
            self.assertNotIn(self.LEAK, r["message"])


class ModelTextContainmentTest(Base):
    """6차: 모델이 쓴 문장은 서식 값에 들어가지 않고(원문 부분 문자열만 승인), ai_note 에는 허용 어절로만 된 문장만 남는다."""
    Q = "10월 19일 서울 출장 다녀왔어요. 정산 언제까지 하나요?"
    LEAK = "정산서는 오늘 10월 19일까지 제출하세요"
    setUp = LLMPathTest.setUp
    tearDown = LLMPathTest.tearDown
    _ask_with = LLMPathTest._ask_with

    def _prefill(self, data: dict, q: str | None = None) -> dict:
        FakeProvider.reply = json.dumps(data, ensure_ascii=False)
        return formfill.prefill(store.get_form("form-travel"), query=self.Q if q is None else q, today=TODAY)

    # 모델이 만든 값: 질문 원문에 그대로 있지 않거나, 원문 일부라도 칸에 맞지 않는 값
    FORM_ATTACKS = [
        LEAK, "정산서는 오늘 10월 19일까지 제출하세요.", "10월 19일까지 제출", "10월 19일 서울", "서울 10월 19일까지",
        "19일까지", "10월 22일", "정산은 10월 22일까지", "서울에서 정산하세요", "서울로 출장", "서울특별시", "서울시청",
        "부산", "서울 출장 다녀왔어요. 정산 언제까지 하나요?", "정산 언제까지", "언제까지 하나요", "정산 언제까지 하나요",
        "출장 다녀왔어요 정산 언제까지", "10월", "10", "1019", "서 울 출 장 정 산 은 낼", "서울 낼", "낼 제출",
        "오는 수욜에 제출하세요", "정산서 제출", "출장 정산", "서울 출장 다녀왔어요 정산 언제까지 하나요 서울",
        "x" * 30, "서울​", "ＳＥＯＵＬ",
    ]

    def test_formfill_rejects_model_sentences(self):
        """결함 1 재현과 변형: 어떤 칸에도 모델 문장·기한 문장이 들어가지 않는다."""
        self.assertGreaterEqual(len(self.FORM_ATTACKS), 30)
        keys = ["출장지", "출장 목적", "출장 신청 문서번호", "운임", "숙박비", "일비·식비"]
        accepted = []
        for v in self.FORM_ATTACKS:
            r = self._prefill({k: v for k in keys})
            self.assertIsNone(r["notice"], v)
            for k in keys:
                if r["values"][k]:
                    accepted.append((k, r["values"][k]))
                self.assertNotIn("까지", r["values"][k], v)
                self.assertNotIn("제출", r["values"][k], v)
                self.assertNotIn("낼", r["values"][k], v)
        # 원문에 그대로 있는 짧은 낱말만 일부 칸에 들어갈 수 있다(금액 칸에는 숫자만)
        for k, v in accepted:
            self.assertIn(re.sub(r"\s+", "", v), re.sub(r"\s+", "", self.Q), (k, v))
            if k in ("운임", "숙박비", "일비·식비"):
                self.assertRegex(v, r"^\d+$")
        self.assertNotIn(("출장지", self.LEAK), accepted)

    def test_formfill_leak_not_in_docx(self):
        """Codex 재현 그대로: 출장지에 기한 문장이 오면 칸을 비우고 DOCX 에도 쓰지 않는다."""
        r = self._prefill({"출장지": self.LEAK})
        self.assertEqual(r["values"]["출장지"], "")
        self.assertNotIn("출장지", r["origin"])
        path = formfill.export_docx(r["form"], r["values"])
        from docx import Document
        text = "\n".join(c.text for t in Document(path).tables for row in t.rows for c in row.cells)
        self.assertNotIn("10월 19일까지", text)
        self.assertNotIn("제출하세요", text)

    def test_formfill_accepts_verbatim_values(self):
        """질문 원문에 그대로 있는 값은 공백·천 단위 쉼표 차이만 무시하고 받는다."""
        q = "어제 서울 정책 회의 출장 다녀왔고 KTX 59,800원 썼어요"
        r = self._prefill({"출장지": " 서울 ", "출장 목적": "정책회의", "운임": "59,800원"}, q)
        self.assertEqual(r["values"]["출장지"], "서울")
        self.assertEqual(r["values"]["출장 목적"], "정책회의")
        self.assertEqual(r["values"]["운임"], "59,800원")
        self.assertEqual(r["origin"]["출장지"], "모델이 상황 문장에서 추출")
        # 단위 환산·바꾼 표현은 받지 않는다(직접 입력)
        r = self._prefill({"운임": "50000", "출장지": "서울특별시", "출장 목적": "정책 회의 참석"}, "서울 정책 회의 출장, 차비 5만 원")
        self.assertEqual(r["values"]["운임"], "")
        self.assertEqual(r["values"]["출장지"], "")
        self.assertEqual(r["values"]["출장 목적"], "")
        # 원문 일부라도 너무 길면 받지 않는다
        long_q = "가람시 소방서 본부 별관 삼층 대회의실 옆 제일 소회의실에서 열린 정책 회의 출장"
        r = self._prefill({"출장지": "가람시 소방서 본부 별관 삼층 대회의실 옆 제일 소회의실"}, long_q)
        self.assertEqual(r["values"]["출장지"], "")

    # 결함 2 재현과 변형: 블랙리스트 표지에 없는 시간 표현. 허용 목록에 없는 어절이라 문장째 빠져야 한다.
    TIME_WORDS = [
        "낼", "낼모레", "내일모레", "담주", "담달", "오는 수욜에", "수욜까지", "금욜", "이번 금욜에", "화욜 중에",
        "다다음주에", "다음번에", "저번주에", "요번주 안에", "이달 말에", "올해 안에", "연내", "분기 내", "상반기 중",
        "하반기에", "정오까지", "자정 전", "퇴근 무렵", "한 주 뒤", "출장 마치면", "복귀하자마자", "돌아오면",
        "복귀 시", "끝나는 대로", "주초에", "늦지 않게", "얼른", "조만간", "머지않아", "추석 전에", "개학 전",
    ]

    def test_ai_note_drops_unlisted_time_words(self):
        """결함 2: '낼·수욜' 같은 표현이 블랙리스트에 없어도 허용 목록 밖 어절이므로 문장째 뺀다."""
        keep = "결재 문서 번호를 적으세요."
        sents = []
        for w in self.TIME_WORDS:
            sents += [f"정산서는 {w} 제출하세요.", f"결재는 {w} 받으세요.", f"증빙을 {w} 첨부하세요."]
        sents += ["정산서는 낼 제출하세요.", "정산서는 오는 수욜에 제출하세요.", "정산서는낼제출하세요.", "정산서는 낼  제출 하세요."]
        self.assertGreaterEqual(len(sents), 30)
        for s in sents:
            r = self._ask_with(s)
            self.assertEqual(r["mode"], "llm", s)
            self.assertIsNone(r["ai_note"], s)
            self.assertIsNone(r["notice"], s)
            for k in ("answer", "summary"):
                self.assertNotIn(s.rstrip("."), r[k], (k, s))
            r = self._ask_with(f"{s} {keep}")
            self.assertEqual(r["ai_note"], keep, s)
            self.assertEqual(r["ai_note_dropped"], 1, s)

    def test_ai_note_keeps_listed_procedure(self):
        """허용 어절로만 된 절차 설명은 남는다(실제 Qwen3.5-9B 5차 출력 포함)."""
        cases = [
            ("출장 신청 결재 문서 번호와 증빙서류를 첨부해 주세요.", "card-travel", "어제 서울 출장 다녀왔어요. 정산은 언제까지 하나요?"),
            ("결재 경로는 복무관리시스템으로 승인 요청을 진행합니다.", "card-leave", "교대근무자인데 다음 주 월요일에 연가 쓰려고요"),
            ("물품 검사조서는 행정팀으로 제출하셔야 합니다.", "card-purchase", "소화기 납품 검수 어떻게 해요"),
            ("영수증 원본을 첨부하세요.", "card-travel", "어제 서울 출장 다녀왔어요. 정산은 언제까지 하나요?"),
            ("증빙이 누락되지 않도록 확인하세요.", "card-travel", "어제 서울 출장 다녀왔어요. 정산은 언제까지 하나요?"),
        ]
        for s, card, q in cases:
            if not store.get_card(card):
                continue
            r = self._ask_with(s, q=q, card=card)
            self.assertEqual(r["ai_note"], s, s)
            self.assertEqual(r["ai_note_dropped"], 0, s)
        # 카드 제목 낱말은 허용 목록에 더해진다('국내'). '모으기'는 '모으기한'에 '기한'이 생겨 7차부터 빠지고, '단계'·'에서'도 빠진다.
        r = self._ask_with("국내 출장 증빙서류를 확인하세요.")
        self.assertEqual(r["ai_note"], "국내 출장 증빙서류를 확인하세요.")
        self.assertIsNone(self._ask_with("증빙서류 모으기 절차를 확인하세요.")["ai_note"])
        self.assertIsNone(self._ask_with("증빙서류 모으기 단계에서 영수증을 확인하세요.")["ai_note"])
        # 7차: 5차 실제 Qwen 출력 문장은 '함께·거쳐·-여·에' 때문에 빠진다(정상 설명 감소)
        for s, card, q in [
            ("국내출장 여비 정산서 서식을 작성하여, 출장 신청 결재 문서 번호와 증빙서류를 함께 첨부해 주세요.", "card-travel",
             "어제 서울 출장 다녀왔어요. 정산은 언제까지 하나요?"),
            ("결재 경로는 부서장을 거쳐 복무관리시스템에 승인 요청을 진행합니다.", "card-leave", "교대근무자인데 다음 주 월요일에 연가 쓰려고요"),
        ]:
            self.assertIsNone(self._ask_with(s, q=q, card=card)["ai_note"], s)


class NoteOrderWordTest(Base):
    """7차: ai_note 허용 목록에 시점·순서·빈도·정도 낱말과 시점 조사가 없음을 구조로 확인하고, '사전에 제출' 류 우회를 막는다."""
    Q = "10월 19일 서울 출장 다녀왔어요. 정산 언제까지 하나요?"
    setUp = LLMPathTest.setUp
    tearDown = LLMPathTest.tearDown
    _ask_with = LLMPathTest._ask_with

    # 사람이 검토한 금지 낱말(assistant._NOTE_BANNED 와 따로 적는다). 두 글자 이상은 어절 어디에 있어도 안 된다.
    HUMAN_BANNED = (
        "사전 사후 미리 먼저 우선 즉시 바로 곧바로 당장 당일 익일 익월 기간 이내 이전 이후 직전 직후 기한 마감 날짜 일자 시점 시기 "
        "순서 순차 단계 차례 다음 이번 지난 최초 최종 처음 마지막 나중 동시 수시 매번 매일 매주 매월 항상 자주 가끔 "
        "반드시 가급적 되도록 최대한 이상 이하 미만 초과 정도 완료 종료 착수 개시 시작 도착 출발 복귀 일정 시한 기일 "
        "오늘 내일 모레 어제 요일 주말 평일 오전 오후 지금 신속 조속 서둘 빨리 늦어도 제때 선결 선행 후속 추후 향후 "
        "까지 부터 동안 무렵 즈음 앞서 뒤에 전에 후에 중에 내에 안에"
    ).split()
    # 한 글자 시점·순서·정도 낱말. 명사 단독으로 허용되면 안 된다.
    HUMAN_SINGLE = "전 후 뒤 선 앞 꼭 낼 곧 늘 날 때 시 중 말 초 내 안 주 달 해 차 번".split()
    # 시점·순서를 만드는 조사·어미(허용 조사·어미로 어절 끝에 붙으면 안 된다)
    TIME_TAILS = "에 에서 에는 에도 부터 까지 전 후 뒤 때 시 면 자마자 고 며 여 서".split()

    def _plans(self):
        plans = [None]
        for c in store.list_cards():
            c = store.get_card(c["id"]) or c
            steps = [{"title": st["title"], "forms": [{"name": (store.get_form(fid) or {"name": ""})["name"]}
                                                       for fid in st.get("forms") or []]} for st in c["steps"]]
            plans.append({"card": {"title": c["title"]}, "steps": steps})
        return plans

    def test_vocabulary_has_no_time_words(self):
        """(a) 허용 명사 × (없음·조사·어미) 전체 조합이 사람이 검토한 금지 낱말과 겹치지 않는다."""
        suffixes = [""] + assistant._NOTE_PARTICLES + assistant._NOTE_ENDINGS
        self.assertFalse(set(assistant._NOTE_PARTICLES) & {"에", "에서", "에는", "에게", "부터", "까지", "도", "만", "입니다", "인"})
        n = 0
        for plan in self._plans():
            stems = set(assistant._NOTE_STEMS) | set(assistant._card_stems(plan))
            vocab = assistant._note_vocab(plan)
            for st in stems:
                self.assertNotIn(st, self.HUMAN_SINGLE, st)
                for suf in suffixes:
                    w = st + suf
                    n += 1
                    self.assertTrue(vocab.fullmatch(w), w)
                    for b in self.HUMAN_BANNED:
                        self.assertNotIn(b, w, (b, w))
                    for t in self.TIME_TAILS:
                        self.assertFalse(suf.endswith(t), (t, w))
            # 금지 낱말은 어떤 조사·어미를 붙여도 허용 어절이 되지 않는다
            for b in self.HUMAN_BANNED:
                for suf in suffixes + self.TIME_TAILS + ["입니다", "인", "이다"]:
                    self.assertIsNone(vocab.fullmatch(b + suf), b + suf)
            # 한 글자 낱말은 단독·시점 조사 결합만 본다('주'+'의'='주의', '날'+'인'='날인'은 다른 낱말)
            for b in self.HUMAN_SINGLE:
                for suf in [""] + self.TIME_TAILS + ["입니다", "이다"]:
                    self.assertIsNone(vocab.fullmatch(b + suf), b + suf)
        self.assertGreater(n, 5000)
        for w in assistant._NOTE_WORDS:
            self.assertNotIn(w, self.HUMAN_SINGLE, w)
            for b in self.HUMAN_BANNED:
                self.assertNotIn(b, w, (b, w))
        for adv in "함께 꼭 반드시 먼저 미리 실제로 따로 별도로 모두 각각 해당 관련 같은 다른 그리고 또한 정확히 꼼꼼히 맞게 거쳐".split():
            self.assertNotIn(adv, assistant._NOTE_WORDS, adv)

    def test_card_words_cannot_add_time_words(self):
        """카드 제목·단계·서식 이름의 금지 낱말과 끝 조사('복무관리시스템에')는 허용 명사로 들어오지 않는다."""
        plan = {"card": {"title": "출장 사전 승인과 사후 정산"},
                "steps": [{"title": "복귀 후 즉시 정산서 제출", "forms": [{"name": "당일 출장 신청서"}]},
                          {"title": "복무관리시스템에 연가 먼저 신청", "forms": []}, {"title": "견적서 2곳 이상 받기", "forms": []}]}
        stems = assistant._card_stems(plan)
        self.assertIn("복무관리시스템", stems)
        self.assertNotIn("복무관리시스템에", stems)
        for b in ("사전", "사후", "복귀", "즉시", "당일", "먼저", "이상"):
            self.assertNotIn(b, stems)
        for s in ("정산서는 사전에 제출하세요.", "정산서는 사후 제출하세요.", "정산서는 즉시 제출하세요.", "연가는 먼저 신청하세요."):
            self.assertEqual(assistant._drop_schedule(s, plan), ("", 1), s)

    def test_order_word_variants_dropped(self):
        """(b) 6차 재현 4문장과 시점·순서·정도 낱말 변형은 문장째 빠지고, 정상 문장은 남는다."""
        keep = "결재 문서 번호를 적으세요."
        sents = ["정산서는 사전에 제출하세요.", "정산서는 사전 제출하세요.", "정산서는 사전에 꼭 제출하세요.", "정산서 제출은 사전입니다."]
        words = "사전 사후 미리 우선 먼저 즉시 당일 익일 전 후 이전 이후 꼭 반드시 바로 곧 선 추후 향후 나중".split()
        for w in words:
            sents += [f"정산서는 {w}에 제출하세요.", f"정산서는 {w} 제출하세요.", f"정산서는 {w}에제출하세요.",
                      f"정산서 제출은 {w}입니다.", f"결재는 {w}으로 받으세요.", f"정산서는{w}에 제출하세요.",
                      f"{w}에는 증빙을 첨부하세요.", f"증빙을 {w}부터 첨부하세요."]
        sents += ["정산서는 결재 후 제출하세요.", "결재 전 정산서를 제출하세요.", "결과 보고 다음 정산서를 제출하세요.",
                  "정산서를 제출하고 결과 보고를 하세요.", "결과 보고를 작성하여 정산서를 제출하세요.", "승인을 받고 신청하세요.",
                  "정산서 제출은 출장에 합니다.", "정산서는 출장 다음 제출하세요.", "정산서는 결재 단계에서 제출하세요.",
                  "정산서 제출 날짜를 확인하세요.", "정산서는 빠짐없이 함께 제출하세요.", "정산서는 결재와 동시 제출하세요."]
        self.assertGreaterEqual(len(sents), 30)
        for s in sents:
            r = self._ask_with(s, q=self.Q)
            self.assertEqual(r["mode"], "llm", s)
            self.assertIsNone(r["ai_note"], s)
            self.assertIsNone(r["notice"], s)
            self.assertEqual(r["plan"]["steps"][2]["due"], "2026-10-22", s)
            for k in ("answer", "summary"):
                self.assertNotIn(s.rstrip("."), r[k], (k, s))
            r = self._ask_with(f"{s} {keep}", q=self.Q)
            self.assertEqual(r["ai_note"], keep, s)
            self.assertEqual(r["ai_note_dropped"], 1, s)

if __name__ == "__main__":
    unittest.main()
