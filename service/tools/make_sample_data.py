"""합성 모의데이터 생성기 (개발용, 심사·구동에는 필요 없음).

가상 기관 '가람시소방서 행정팀'의 규정·지침·서식·메모를 만든다. 모든 기관명·조문·금액·연락처는 지어낸 것이며
실제 법령·소방 내부 문서와 무관하다. 생성물은 sample_data/규정자료/ 에 저장되고, 저장소에는 생성된 파일이 포함돼 있다.

필요: python-docx, openpyxl, reportlab(PDF용, 개발 시에만), lxml
HWPX는 HWPX_SKILL_DIR 환경변수가 가리키는 hwpx 빌드 도구(build_hwpx.py, gonmun 템플릿)로 조립한다.
"""
from __future__ import annotations

import os
import subprocess
import sys
import tempfile
from pathlib import Path
from xml.sax.saxutils import escape

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "sample_data" / "규정자료"
SKILL = Path(os.environ.get("HWPX_SKILL_DIR", ""))

# ------------------------------------------------------------------ 본문
TRAVEL_RULE = ("가람시 공무원 여비 규정(가상)", [
    ("h", "가람시 공무원 여비 규정(가상)"),
    ("p", "※ 이 문서는 공모전 시연을 위해 지어낸 가상 규정으로, 실제 법령·기관 규정과 무관합니다."),
    ("a", "제1조(목적) 이 규정은 가람시 소속 공무원이 공무로 여행하는 경우 지급하는 여비에 관하여 필요한 사항을 정한다."),
    ("a", "제2조(정의) 이 규정에서 \"출장\"이란 공무를 위하여 근무지를 떠나 여행하는 것을 말하고, \"정산\"이란 출장을 마친 뒤 실제 사용한 여비를 확인하여 차액을 지급하거나 반납받는 것을 말한다."),
    ("a", "제5조(출장 신청) ① 출장하려는 공무원은 출발 전날까지 출장 신청서를 작성하여 소속 부서장의 승인을 받아야 한다."),
    ("p", "② 재난 대응 등 긴급한 경우에는 출장 당일에 승인을 받을 수 있으며, 이 경우 복귀 후 사유를 기록한다."),
    ("a", "제7조(운임) ① 철도·버스·항공 운임은 실제 든 비용으로 지급하며, 승차권 또는 영수증을 증빙으로 첨부한다."),
    ("p", "② 자가용 차량을 이용한 경우에는 출발지와 도착지 사이 거리에 1킬로미터당 200원(가상 기준)을 곱한 금액을 지급한다."),
    ("a", "제8조(숙박비) 숙박비는 1박당 7만원(가상 기준) 한도에서 실제 든 비용을 지급하며, 숙박 영수증을 첨부한다. 공용 시설에 숙박한 경우에는 지급하지 않는다."),
    ("a", "제9조(일비와 식비) 일비와 식비는 각각 1일 2만5천원(가상 기준)을 정액으로 지급한다. 다만, 식사가 제공된 경우에는 식비를 줄여 지급한다."),
    ("a", "제12조(여비의 정산) ① 출장을 마친 공무원은 출장이 끝난 날부터 7일 이내에 여비 정산서에 증빙서류를 첨부하여 여비 담당자에게 제출하여야 한다."),
    ("p", "② 여비 담당자는 정산서를 받은 날부터 10일 이내에 정산 결과에 따라 차액을 지급하거나 반납을 통지한다."),
    ("p", "③ 증빙서류를 잃어버린 경우에는 카드 사용 내역 등 사용 사실을 확인할 수 있는 자료로 대신할 수 있다."),
    ("a", "제13조(출장 결과 보고) 출장을 마친 공무원은 복귀한 날 출장 결과를 소속 부서장에게 보고하여야 한다. 다만, 단순 서류 전달 등 경미한 출장은 보고를 생략할 수 있다."),
])

AGENCY_RULE = ("가람시소방서 행정업무 처리 기준(가상)", [
    ("h", "가람시소방서 행정업무 처리 기준(가상)"),
    ("p", "※ 공모전 시연용 가상 기관 기준입니다. 공통 규정을 보충하며, 둘이 다르면 이 기준을 먼저 따릅니다."),
    ("a", "제1조(목적) 이 기준은 가람시소방서 직원의 출장 정산, 물품 검수, 행사 준비 등 서무 업무의 처리 기한과 절차를 정한다."),
    ("a", "제3조(출장 여비 정산 기한) ① 직원은 「가람시 공무원 여비 규정(가상)」 제12조에도 불구하고 출장이 끝난 날의 다음 업무일부터 3업무일 이내에 여비 정산서를 행정팀에 제출한다."),
    ("p", "② 12월에 끝난 출장의 정산서는 회계연도 마감을 위해 12월 26일까지 제출한다."),
    ("p", "③ 정산서에는 출장 신청 결재 문서 번호를 함께 적는다."),
    ("a", "제4조(물품 검수) 물품을 납품받은 부서는 납품일부터 5업무일 이내에 검수를 마치고 물품 검사조서를 행정팀에 제출한다."),
    ("a", "제5조(행사 준비 일정) ① 대외 행사를 여는 부서는 행사일 21일 전까지 행사 계획을 결재받는다."),
    ("p", "② 행사에 필요한 물품은 행사일 14일 전까지 구매를 요구한다."),
    ("p", "③ 참석 안내문은 행사일 7일 전까지 발송한다."),
    ("p", "④ 행사가 끝나면 행사일 다음 업무일부터 5업무일 이내에 결과를 보고한다."),
    ("a", "제6조(서식) 이 기준에 따른 서류는 행정팀 공유 폴더의 기관 서식을 사용한다."),
])

PURCHASE_GUIDE = ("물품 구매 및 계약 사무 처리 지침(가상)", [
    ("h", "물품 구매 및 계약 사무 처리 지침(가상)"),
    ("p", "※ 공모전 시연용 가상 지침입니다. 금액 기준과 기한은 실제 계약 법령과 다릅니다."),
    ("a", "제1조(목적) 이 지침은 가람시 소속 기관의 소액 물품 구매 절차를 정한다."),
    ("a", "제3조(구매 요구) 물품이 필요한 부서는 사용 예정일 14일 전까지 물품 구매 요구서를 작성하여 계약 담당자에게 제출한다."),
    ("a", "제4조(견적) 추정가격 2천만원(가상 기준) 이하의 물품은 2인 이상에게 견적서를 받아 가장 낮은 금액을 제시한 자와 계약할 수 있다."),
    ("a", "제6조(검사) ① 계약 담당자는 납품을 받은 날부터 14일 이내에 검사를 마쳐야 한다."),
    ("p", "② 검사를 마치면 물품 검사조서를 작성하고 검사자가 서명한다."),
    ("a", "제7조(대금 지급) 검사를 마친 뒤 대금 청구를 받으면 청구를 받은 날부터 5일 이내에 대금을 지급한다."),
])

LEAVE_GUIDE = [
    "연가·복무 처리 안내(가상)",
    "※ 공모전 시연용 가상 안내문입니다.",
    "제1조(연가 신청) 연가를 쓰려는 직원은 사용하려는 날의 전 업무일까지 복무관리시스템에 신청하고 부서장의 승인을 받는다.",
    "제2조(긴급한 연가) 질병·사고 등으로 미리 신청하지 못한 경우에는 당일 오전 중 부서장에게 알리고, 복귀한 날 신청한다.",
    "제3조(교대근무자) 교대근무 직원은 근무표 조정을 위해 사용하려는 날 7일 전까지 신청한다.",
]

MEMO = """서무담당 업무 메모 (개인, 가상)
KTX 승차권 영수증은 코레일 앱 '승차권 확인'에서 다시 내려받을 수 있음.
여비 정산서 낼 때 출장 신청 결재 문서 번호를 꼭 같이 적기 (작년에 반려된 적 있음).
숙박비 한도 넘으면 초과분은 본인 부담. 영수증에 숙박일수 나오는지 확인.
행사 물품은 납품 후 검수까지 5업무일이라 행사 직전 납품은 피하는 게 좋음.
12월 출장은 26일 마감이니 크리스마스 전에 미리 정리.
"""

DEADLINE_TABLE = [
    ["업무", "처리 기한", "근거", "제출처"],
    ["출장 결과 보고", "복귀한 날", "여비 규정(가상) 제13조", "부서장"],
    ["여비 정산서 제출", "출장 종료 다음 업무일부터 3업무일 이내", "행정업무 처리 기준(가상) 제3조", "행정팀"],
    ["물품 검수", "납품일부터 5업무일 이내", "행정업무 처리 기준(가상) 제4조", "행정팀"],
    ["행사 계획 결재", "행사일 21일 전", "행정업무 처리 기준(가상) 제5조", "부서장"],
    ["행사 결과 보고", "행사 다음 업무일부터 5업무일 이내", "행정업무 처리 기준(가상) 제5조", "부서장"],
    ["연가 신청", "사용 전 업무일까지", "연가·복무 처리 안내(가상) 제1조", "복무관리시스템"],
]

FORM_TRAVEL = ("국내출장 여비 정산서(기관 서식, 가상)", [
    ("소속", ""), ("직급", ""), ("성명", ""), ("출장 기간", ""), ("출장지", ""), ("출장 목적", ""),
    ("출장 신청 문서번호", ""), ("운임", "원"), ("숙박비", "원"), ("일비·식비", "원"), ("합계", "원"),
    ("작성일", ""),
])
FORM_INSPECT = ("물품 검사조서(기관 서식, 가상)", [
    ("품명", ""), ("규격·수량", ""), ("계약 금액", "원"), ("납품일", ""), ("검사일", ""),
    ("검사 결과", "합격 / 불합격"), ("검사자 소속·성명", ""),
])


# ------------------------------------------------------------------ HWPX
FIRST_P = None


def _first_paragraph() -> str:
    """gonmun 템플릿 section0.xml 의 첫 문단(secPr 포함)을 그대로 쓴다."""
    global FIRST_P
    if FIRST_P is None:
        src = (SKILL / "templates" / "gonmun" / "Contents" / "section0.xml")
        if not src.exists():
            src = SKILL / "templates" / "base" / "Contents" / "section0.xml"
        s = src.read_text(encoding="utf-8")
        start = s.find("<hp:p ")
        FIRST_P = s[start:s.find("</hp:p>", start) + len("</hp:p>")]
    return FIRST_P


NS = ('xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph" xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section" '
      'xmlns:hc="http://www.hancom.co.kr/hwpml/2011/core" xmlns:hh="http://www.hancom.co.kr/hwpml/2011/head"')
_pid = [1000000100]


def _p(text: str, char: int = 0, para: int = 0) -> str:
    _pid[0] += 1
    return (f'<hp:p id="{_pid[0]}" paraPrIDRef="{para}" styleIDRef="0" pageBreak="0" columnBreak="0" merged="0">'
            f'<hp:run charPrIDRef="{char}"><hp:t>{escape(text)}</hp:t></hp:run></hp:p>')


def _cell(text: str, col: int, row: int, width: int, header: bool) -> str:
    _pid[0] += 1
    run = f'<hp:t>{escape(text)}</hp:t>' if text else "<hp:t/>"
    return (f'<hp:tc name="" header="0" hasMargin="0" protect="0" editable="0" dirty="1" borderFillIDRef="{4 if header else 3}">'
            '<hp:subList id="" textDirection="HORIZONTAL" lineWrap="BREAK" vertAlign="CENTER" linkListIDRef="0" '
            'linkListNextIDRef="0" textWidth="0" textHeight="0" hasTextRef="0" hasNumRef="0">'
            f'<hp:p id="{_pid[0]}" paraPrIDRef="22" styleIDRef="0" pageBreak="0" columnBreak="0" merged="0">'
            f'<hp:run charPrIDRef="{10 if header else 0}">{run}</hp:run></hp:p></hp:subList>'
            f'<hp:cellAddr colAddr="{col}" rowAddr="{row}"/><hp:cellSpan colSpan="1" rowSpan="1"/>'
            f'<hp:cellSz width="{width}" height="2400"/><hp:cellMargin left="510" right="510" top="141" bottom="141"/></hp:tc>')


def _table(rows: list[tuple[str, str]]) -> str:
    _pid[0] += 1
    tid = _pid[0]
    w0, w1 = 12756, 29764
    trs = "".join(f"<hp:tr>{_cell(k, 0, i, w0, True)}{_cell(v, 1, i, w1, False)}</hp:tr>" for i, (k, v) in enumerate(rows))
    _pid[0] += 1
    return (f'<hp:p id="{_pid[0]}" paraPrIDRef="0" styleIDRef="0" pageBreak="0" columnBreak="0" merged="0"><hp:run charPrIDRef="0">'
            f'<hp:tbl id="{tid}" zOrder="0" numberingType="TABLE" textWrap="TOP_AND_BOTTOM" textFlow="BOTH_SIDES" lock="0" '
            f'dropcapstyle="None" pageBreak="CELL" repeatHeader="0" rowCnt="{len(rows)}" colCnt="2" cellSpacing="0" borderFillIDRef="3" noAdjust="0">'
            f'<hp:sz width="42520" widthRelTo="ABSOLUTE" height="{2400 * len(rows)}" heightRelTo="ABSOLUTE" protect="0"/>'
            '<hp:pos treatAsChar="1" affectLSpacing="0" flowWithText="1" allowOverlap="0" holdAnchorAndSO="0" vertRelTo="PARA" '
            'horzRelTo="COLUMN" vertAlign="TOP" horzAlign="LEFT" vertOffset="0" horzOffset="0"/>'
            '<hp:outMargin left="0" right="0" top="0" bottom="0"/><hp:inMargin left="0" right="0" top="0" bottom="0"/>'
            f'{trs}</hp:tbl></hp:run></hp:p>')


def build_hwpx(name: str, body_xml: str, title: str) -> None:
    section = f"<?xml version='1.0' encoding='UTF-8'?>\n<hs:sec {NS}>{_first_paragraph()}{body_xml}</hs:sec>"
    with tempfile.NamedTemporaryFile("w", suffix=".xml", delete=False, encoding="utf-8") as f:
        f.write(section)
    out = OUT / name
    subprocess.run([sys.executable, str(SKILL / "scripts" / "build_hwpx.py"), "--template", "gonmun", "--section", f.name,
                    "--title", title, "--creator", "가람시소방서(가상)", "--output", str(out)], check=True)
    subprocess.run([sys.executable, str(SKILL / "scripts" / "validate.py"), str(out)], check=True)
    os.unlink(f.name)


def rule_hwpx(name: str, spec: tuple[str, list]) -> None:
    title, items = spec
    xml = ""
    for kind, text in items:
        xml += _p(text, 7, 20) if kind == "h" else _p(text, 0, 0)
    build_hwpx(name, xml, title)


def form_hwpx(name: str, spec: tuple[str, list]) -> None:
    title, rows = spec
    xml = _p(title, 7, 20) + _p("") + _table(rows) + _p("") + _p("위와 같이 작성합니다.       작성자:            (서명)")
    build_hwpx(name, xml, title)


# ------------------------------------------------------------------ PDF / DOCX / XLSX / TXT
def build_pdf(name: str, spec: tuple[str, list]) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    font = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / "malgun.ttf"
    pdfmetrics.registerFont(TTFont("Malgun", str(font)))
    body = ParagraphStyle("b", fontName="Malgun", fontSize=10.5, leading=17, spaceAfter=8)
    head = ParagraphStyle("h", fontName="Malgun", fontSize=17, leading=24, alignment=1, spaceAfter=14)
    flow = []
    for kind, text in spec[1]:
        flow.append(Paragraph(escape(text), head if kind == "h" else body))
    flow.insert(1, Spacer(1, 6))
    SimpleDocTemplate(str(OUT / name), pagesize=A4, title=spec[0], author="가람시(가상)").build(flow)


def build_docx(name: str, lines: list[str]) -> None:
    from docx import Document
    from docx.shared import Pt
    d = Document()
    for i, ln in enumerate(lines):
        p = d.add_paragraph()
        r = p.add_run(ln)
        if i == 0:
            r.bold = True
            r.font.size = Pt(16)
    d.save(OUT / name)


def build_xlsx(name: str) -> None:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "처리기한"
    for r in DEADLINE_TABLE:
        ws.append(r)
    for col, w in zip("ABCD", (18, 36, 34, 16)):
        ws.column_dimensions[col].width = w
    wb.save(OUT / name)


def main() -> None:
    if not (SKILL / "scripts" / "build_hwpx.py").exists():
        sys.exit("HWPX_SKILL_DIR 환경변수에 hwpx 빌드 도구 경로를 지정하세요.")
    OUT.mkdir(parents=True, exist_ok=True)
    rule_hwpx("01_공통_여비규정_가상.hwpx", TRAVEL_RULE)
    build_pdf("02_공통_물품구매지침_가상.pdf", PURCHASE_GUIDE)
    build_docx("03_공통_연가복무안내_가상.docx", LEAVE_GUIDE)
    rule_hwpx("04_기관_행정업무처리기준_가상.hwpx", AGENCY_RULE)
    build_xlsx("05_기관_업무별처리기한표_가상.xlsx")
    (OUT / "06_개인_서무담당메모_가상.txt").write_text(MEMO, encoding="utf-8")
    form_hwpx("07_서식_국내출장여비정산서_가상.hwpx", FORM_TRAVEL)
    form_hwpx("08_서식_물품검사조서_가상.hwpx", FORM_INSPECT)
    print("생성 완료:", OUT)


if __name__ == "__main__":
    main()
