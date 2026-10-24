"""Exercise running apps using explicitly named synthetic jobs and documents only."""
import argparse
import io
import json
import zipfile
from pathlib import Path

import httpx
from openpyxl import load_workbook
from app import adhoc


PREFIX='[합성 시연 v36]'


def check(port, key=None):
    client=httpx.Client(base_url=f'http://127.0.0.1:{port}',timeout=20)
    def request(method,path,body=None,status=200,**kwargs):
        response=client.request(method,path,json=body,**kwargs)
        assert response.status_code==status,(method,path,response.status_code,response.text[:400])
        return response
    def read():return request('GET',f'/api/adhoc/jobs/{key}').json()
    settings=dict(situation='합성 자료로 교육 참석 희망자 취합 과정을 검증해요.',title=PREFIX+' 교육 희망자 취합',
                  course='공공 AI 실무교육',event_date='2026-11-06',deadline='2026-10-30',departments=['총무팀','행정팀','안전팀'])
    if key is None:key=request('POST','/api/adhoc/jobs',settings,status=201).json()['id']
    j=read();assert j['settings']['title'].startswith(PREFIX),'Only explicitly synthetic jobs may be changed'
    assert not j['files'],'Use a fresh synthetic job'
    def command(action,step=None,status=200,**data):
        nonlocal j
        response=request('PATCH',f'/api/adhoc/jobs/{key}/workflow',dict(action=action,step_id=step,revision=j['workflow']['revision'],data_version=j['version'],**data),status)
        if status==200:j=response.json()
        return response
    def complete(step):command('complete',step,confirmed=True)
    command('complete','review',status=422,confirmed=True)
    command('archive','intake',status=422)
    if not any(s['kind']=='custom' for s in j['workflow']['steps']):command('add',title='부서장 확인',checks=['합성 명단과 검토 사유를 확인했어요.'],before_id='report')
    custom=next(s['id'] for s in j['workflow']['steps'] if s['kind']=='custom')
    command('archive',custom);command('restore',custom)

    title=PREFIX+' 교육 취합 안내 · 공식 근거 아님'
    source=request('POST','/api/documents',status=200,data=dict(title=title,kind='매뉴얼',layer='기관',note='합성 실무 점검 자료입니다. 공식 법령·기관 지침·실제 경험이 아닙니다.'),files={'file':('synthetic-education.txt','교육 취합 안내\n교육명·교육일·회신 마감·대상 부서를 확인합니다.\n희망자가 없는 부서도 회신 여부를 기록합니다.'.encode('utf-8'),'text/plain')}).json()
    def select_reference():
        refs=request('GET',f'/api/adhoc/jobs/{key}/workflow/references',params={'q':title}).json()
        ref=next(r for r in refs if r['origin']=='document' and r['title']==title)
        command('edit','request',references=[{k:ref[k] for k in ('origin','id','unit','sha256')}])
    select_reference()
    for step in ('intake','request'):complete(step)
    response=request('GET',f'/api/adhoc/jobs/{key}/demo.zip')
    archive=zipfile.ZipFile(io.BytesIO(response.content))
    for i,name in enumerate(sorted(n for n in archive.namelist() if n.endswith('.xlsx') and not n.startswith('99_'))):
        if port==8767 and i==2:
            j=request('POST',f'/api/adhoc/jobs/{key}/empty-replies',dict(version=j['version'],department=j['settings']['departments'][i],reason='합성 시연에서 희망자 없음 회신을 확인함')).json();continue
        j=request('POST',f'/api/adhoc/jobs/{key}/files',status=201,data={'version':str(j['version']),'department':j['settings']['departments'][i]},files={'file':(name,archive.read(name),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}).json()
    request('POST',f'/api/adhoc/jobs/{key}/files',status=422,data={'version':str(j['version']),'department':'총무팀'},files={'file':('wrong.xlsx',archive.read('99_다른서식_거절예시.xlsx'),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert read()==j
    duplicates=[r for r in j['rows'] if any('중복' in issue for issue in r['issues'])]
    assert len(duplicates)==2
    row=duplicates[-1]
    j=request('PATCH',f"/api/adhoc/jobs/{key}/rows/{row['id']}",dict(version=j['version'],values=row['values'],decision='exclude',reason='합성 예제의 같은 회신 중복을 원본과 대조하여 제외')).json()
    missing=next(r for r in j['rows'] if not r['values'][1])
    values=missing['values'][:];values[1]='모의보완직원'
    j=request('PATCH',f"/api/adhoc/jobs/{key}/rows/{missing['id']}",dict(version=j['version'],values=values,decision='include',reason='합성 원본에서 누락된 성명을 확인하여 보완')).json()
    j=request('POST',f'/api/adhoc/jobs/{key}/review-clean',dict(version=j['version'],reason='합성 원본과 대조하여 정상 회신을 포함함')).json()
    assert j['can_finalize']
    for step in ('collect','review',custom):complete(step)
    j=request('POST',f'/api/adhoc/jobs/{key}/finalize',dict(version=j['version']),status=201).json();complete('report')
    version=j['version'];export=j['exports'][0]
    assert j['workflow']['completed_count']==6 and export['is_current']
    command('move',custom,before_id='report')
    assert j['version']==version and j['exports'][0]['is_current']
    complete_snapshot=j

    request('POST',f"/api/documents/{source['id']}/file",files={'file':('synthetic-education-v2.txt','교육 취합 안내\n교육명·교육일·회신 마감·대상 부서를 확인합니다.\n마감 직전 미회신 부서를 다시 확인합니다.'.encode('utf-8'),'text/plain')})
    j=read();source_changed=next(s for s in j['workflow']['steps'] if s['id']=='request')
    assert source_changed['status']=='needs_review' and not source_changed['references'][0]['current']
    command('complete','request',status=422,confirmed=True)
    select_reference();complete('request')

    raw=adhoc.workbook(j['settings'],[['총무팀','모의추가검증자','주무관',j['settings']['course'],j['settings']['event_date'],'희망']])
    j=request('POST',f'/api/adhoc/jobs/{key}/files',status=201,data={'version':str(j['version']),'department':'총무팀'},files={'file':('synthetic-additional.xlsx',raw,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}).json()
    changed_snapshot=j
    for step in ('collect','review','report',custom):assert next(s for s in j['workflow']['steps'] if s['id']==step)['status']=='needs_review'
    assert not j['exports'][0]['is_current']
    command('complete','review',status=422,confirmed=True)
    j=request('POST',f'/api/adhoc/jobs/{key}/review-clean',dict(version=j['version'],reason='합성 추가 회신을 원본과 대조하여 포함')).json()
    j=request('POST',f'/api/adhoc/jobs/{key}/finalize',dict(version=j['version']),status=201).json()
    for step in ('collect','review',custom,'report'):complete(step)
    assert j['workflow']['completed_count']==6
    final_export=next(e for e in j['exports'] if e['is_current'])
    response=request('GET',f"/api/adhoc/jobs/{key}/exports/{final_export['id']}")
    wb=load_workbook(io.BytesIO(response.content));assert wb['제출용 명단'].max_row>1;wb.close()
    assert read()==j
    client.close()
    return dict(port=port,job_id=key,url=f'http://127.0.0.1:{port}/adhoc.html?v=36#job-{key}',passed=True,
                complete_snapshot=complete_snapshot,source_changed=source_changed,data_changed_snapshot=changed_snapshot,final=j)


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--main-job');args=parser.parse_args()
    results=[check(8766,args.main_job),check(8767)]
    target=Path(__file__).resolve().parents[2]/'work/education-workflow-v36-live.json'
    target.write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps([dict(port=r['port'],job_id=r['job_id'],passed=r['passed'],completed=r['final']['workflow']['completed_count'],data_version=r['final']['version']) for r in results],ensure_ascii=False))
